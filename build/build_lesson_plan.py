#!/usr/bin/env python3
"""Build the WSQ Lesson Plan (.docx) for TGS-2025054479.

Slide references are read from build/slide_map.json, which build_slides.py writes,
so the schedule can never cite a slide number the deck does not have.

Hours: 14 h facilitation + 2 h assessment = 16 h over 2 days (per the LMS record).
Day 1 = 7 h facilitation. Day 2 = 7 h facilitation + 2 h assessment.
"""
import os, sys, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT

import course_data as C
import prodoc as P

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, ".."))
ASSETS = os.path.join(REPO, "courseware", "assets")
LOGO = os.path.join(ASSETS, "tertiary-infotech-logo.png")

try:
    SM = json.load(open(os.path.join(HERE, "slide_map.json")))
except OSError:
    SM = {}


def sref(a, b=None):
    """Format a slide reference from the generated slide map."""
    if a not in SM:
        return "—"
    start = SM[a]
    if b and b in SM:
        end = SM[b] - 1
        return f"Slides {start}–{end}" if end > start else f"Slide {start}"
    return f"Slide {start}"


VERSIONS = [
    ("1.0", "19 Apr 2025", "First version", "Tertiary Infotech Academy"),
    ("2.0", "8 May 2025", "Second version", "Tertiary Infotech Academy"),
    ("3.0", "—", "New format", "Tertiary Infotech Academy"),
    ("4.0", "17 Oct 2025", "Update company name", "Tertiary Infotech Academy"),
    ("5.0", C.VERSION_DATE,
     "Full content rebuild against ISO 14001:2015 and ISO 19011:2018. Schedule corrected "
     "to 14 h facilitation + 2 h assessment with no timing gaps. Three case-study "
     "activities added (GreenPack, AquaPure, EcoTech) with slide references generated "
     "from the deck.", "Tertiary Infotech Academy"),
    ("6.0", C.VERSION_DATE,
     "Added the WSQ funding validity period (11 Apr 2025 to 10 Apr 2027) to the course "
     "information, the funding-criteria slide and the Learner Guide.",
     "Tertiary Infotech Academy"),
]

def _lay(start_hhmm, blocks):
    """Lay blocks end-to-end from a start time. blocks = (minutes, topic, method, res).
    Returns the (start, end, duration, topic, method, resources) rows the table needs."""
    h, m = map(int, start_hhmm.split(":"))
    t = h * 60 + m
    rows = []
    for mins, topic, method, res in blocks:
        s0 = f"{t // 60}:{t % 60:02d}"
        t += mins
        s1 = f"{t // 60}:{t % 60:02d}"
        if mins >= 60:
            dur = f"{mins // 60} h" + (f" {mins % 60} min" if mins % 60 else "")
        else:
            dur = f"{mins} min"
        rows.append((s0, s1, dur, topic, method, res))
    return rows


def _ampm(hhmm):
    """'17:45' -> '5:45 pm'. The schedule tables stay 24-hour; prose uses am/pm."""
    h, m = map(int, hhmm.split(":"))
    suffix = "am" if h < 12 else "pm"
    hh = h if 1 <= h <= 12 else (h - 12 if h > 12 else 12)
    return f"{hh}:{m:02d} {suffix}"


def _hours(rows):
    """(facilitation, assessment) minutes — breaks excluded."""
    fac = ass = 0
    for s0, s1, _d, topic, *_ in rows:
        a = [int(x) for x in s0.split(":")]; b = [int(x) for x in s1.split(":")]
        mins = (b[0] * 60 + b[1]) - (a[0] * 60 + a[1])
        if topic.strip() in ("Tea Break", "Lunch Break"):
            continue
        if "FINAL ASSESSMENT" in topic:
            ass += mins
        else:
            fac += mins
    return fac, ass


_D1 = [
    (15, "Digital Attendance (AM) · Welcome and course introduction\n"
         "Trainer introduction · Learner introduction and ice-breaker\n"
         "Overview of course structure · Ground rules · Learning outcomes",
     "Admin · Ice-breaker",
     f"AM Digital Attendance (SSG QR) · Attendance sheet · {sref('admin','topic1')}"),
    (90, "LU1: Introduction to ISO 14001:2015\n"
         "T1: ISO 14001:2015 Internal Auditor Framework (K1, K2, A1)\n"
         "• Overview of the International Professional Practices Framework (IPPF)\n"
         "• Definition of Internal Auditing · Code of Ethics · Core Principles · Standards\n"
         "• Independence and objectivity in practice",
     "Lecture · Discussion",
     f"Trainer slides {sref('topic1')}+ · Learner Guide Topic 1 · TV · Whiteboard · Wi-Fi"),
    (15, "Tea Break", "—", "—"),
    (105, "T1 continued (K1, K2)\n"
          "• Introduction to ISO 14001:2015 and the Environmental Management System\n"
          "• The three intended outcomes · PDCA and the Annex SL high level structure\n"
          "• Clause-by-clause walkthrough: Clauses 4, 5 and 6\n"
          "• Environmental aspects and impacts · compliance obligations (6.1.3)",
     "Lecture · Worked examples · Peer sharing",
     "Trainer slides · Learner Guide Topic 1 · Clause map handout · Whiteboard"),
    (45, "Lunch Break", "—", "—"),
    (90, "Digital Attendance (PM)\n"
          "T1 continued — Apply ISO 14001:2015 to the Organisation (K2, A1)\n"
          "• Risks and opportunities vs environmental aspects · objectives (6.2)\n"
          "• Clauses 7, 8, 9 and 10 · documented information (maintain vs retain)\n"
          "ACTIVITY 1 — GreenPack: framework, context and scope (45 min)",
     "Lecture · Case study · Group discussion",
     f"PM Digital Attendance · Activity 1 brief (activities/) · {sref('act1','topic2')}"),
    (15, "Tea Break", "—", "—"),
    (105, "LU2: Auditing Methodologies\n"
          "T2: Auditing ISO 14001:2015 (K3, A2)\n"
          "• ISO 19011:2018 — guidance vs requirement · the seven audit principles\n"
          "• Managing the audit programme (Clause 5) vs conducting an audit (Clause 6)\n"
          "• Internal audit methodologies and procedures of the organisation",
     "Lecture · Discussion",
     f"Trainer slides {sref('topic2')}+ · Learner Guide Topic 2 · TV · Whiteboard"),
    (15, "Day 1 recap · Q&A · preview of Day 2", "Review · Q&A",
     "Trainer slides · Learner Guide"),
]

_D2 = [
    (15, "Digital Attendance (AM) · Day 1 recap", "Admin · Review",
     "AM Digital Attendance (SSG QR) · Attendance sheet"),
    (90, "T2 continued — Phases of Auditing (K3, A2)\n"
          "• The six-stage audit lifecycle (ISO 19011 §6.2–6.7)\n"
          "• Initiating · preparing · document review · the audit plan\n"
          "• Building a risk-based audit checklist · sampling and the audit trail\n"
          "• Audit evidence: the four sources · sufficient and appropriate evidence",
     "Lecture · Worked examples",
     "Trainer slides · Learner Guide Topic 2 · Checklist template · Whiteboard"),
    (15, "Tea Break", "—", "—"),
    (90, "T2 continued — Engagement with Stakeholders (A2)\n"
         "• Questioning technique: the funnel · 5W1H · 'show me' · PLOD\n"
         "• Opening and closing meeting agendas\n"
         "• Engaging operations staff, management, clients and regulators\n"
         "ACTIVITY 2 — AquaPure: audit plan, checklist and interviews (60 min)",
     "Case study · Role play · Group discussion",
     f"Activity 2 brief (activities/) · {sref('act2','topic3')} · Flip chart"),
    (45, "Lunch Break", "—", "—"),
    (105, "Digital Attendance (PM)\n"
          "LU3: Risk Assessment Techniques\n"
          "T3: Risk Assessment (K4, A3)\n"
          "• Conducting risk assessment · significance evaluation methods\n"
          "• The 5×5 matrix · multi-criteria scoring · FMEA (S × O × D)\n"
          "• Override rules · auditing the significance method",
     "Lecture · Worked examples · Discussion",
     f"PM Digital Attendance · Trainer slides {sref('topic3')}+ · Risk matrix handout"),
    (15, "Tea Break", "—", "—"),
    (105, "T3 continued — Corrective Action and Reporting (K4, A3)\n"
          "• Writing nonconformity statements: requirement + evidence + gap\n"
          "• Grading findings · root cause analysis (5 Whys)\n"
          "• Correction vs corrective action · writing effective CARs\n"
          "• Preparing the audit report · the sampling statement\n"
          "ACTIVITY 3 — EcoTech: risk assessment, CARs and the report (75 min)",
     "Case study · Group work · Presentation",
     f"Activity 3 brief (activities/) · {sref('act3','close')} · CAR form · Flip chart"),
    (15, "Course summary · Briefing for Assessment · Assessment Digital Attendance · "
         "TRAQOM survey",
     "Admin · Briefing",
     f"Assessment Digital Attendance · TRAQOM (LMS) · {sref('assessment','end')}"),
    (60, "FINAL ASSESSMENT — Written Assessment, Short Answer Questions (WA/SAQ)\n"
         "4 questions covering K1, K2, K3 and K4 · open book · individual",
     "Assessment",
     "WA (SAQ) question paper · Assessment Plan · Assessment Summary Record"),
    (60, "FINAL ASSESSMENT — Case Study (CS)\n"
         "3 scenario-based questions covering A1, A2 and A3 · open book · individual",
     "Assessment",
     "CS question paper · Assessment Plan · Assessment Summary Record"),
]

DAY1 = _lay("9:30", _D1)
DAY2 = _lay("9:30", _D2)

_f1, _a1 = _hours(DAY1)
_f2, _a2 = _hours(DAY2)
assert _f1 + _f2 == C.TRAINING_HOURS * 60, (
    f"facilitation is {_f1 + _f2} min, expected {C.TRAINING_HOURS * 60}")
assert _a1 + _a2 == C.ASSESSMENT_HOURS * 60, (
    f"assessment is {_a1 + _a2} min, expected {C.ASSESSMENT_HOURS * 60}")


def hdr(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = "Arial"
    return h


def para(doc, text, size=11, bold=False, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    P.md_runs(p, text, size=size)
    for r in p.runs:
        if bold:
            r.bold = True
    return p


def schedule_table(doc, rows):
    t = doc.add_table(rows=0, cols=6)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdrs = ["Start", "End", "Duration", "Topics / Content / Activity",
            "Instructional Methods", "Resources"]
    widths = [Cm(1.5), Cm(1.5), Cm(2.0), Cm(11.5), Cm(4.0), Cm(5.5)]
    hr = t.add_row().cells
    for i, h in enumerate(hdrs):
        hr[i].text = ""
        r = hr[i].paragraphs[0].add_run(h)
        r.bold = True; r.font.size = Pt(9); r.font.name = "Arial"
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        P.shade_cell(hr[i], "1F6FEB")
    for row in rows:
        cells = t.add_row().cells
        is_break = row[3].strip() in ("Tea Break", "Lunch Break")
        for i, val in enumerate(row):
            cells[i].text = ""
            for j, line in enumerate(str(val).split("\n")):
                p = cells[i].paragraphs[0] if j == 0 else cells[i].add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(line)
                r.font.size = Pt(8.5); r.font.name = "Arial"
                r.bold = is_break or line.startswith(("LU", "T1:", "T2:", "T3:",
                                                      "ACTIVITY", "FINAL"))
            if is_break:
                P.shade_cell(cells[i], "FFF7E6")
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = widths[i]
    return t


def main():
    doc = Document()
    # landscape so the 6-column schedule is readable
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(1.5)

    P.style_base(doc)
    P.add_cover_page(doc, "Lesson Plan", C.TITLE, "6.0", C.COURSE_CODE, org_logo=LOGO)
    P.add_version_control(doc, VERSIONS)
    P.add_toc(doc)

    hdr(doc, "1. Course Information", 1)
    info = [("Course Title", C.TITLE),
            ("TGS Reference Number", C.COURSE_CODE),
            ("Training Provider", f"{C.ORG} ({C.UEN})"),
            ("TSC Title / Code", f"{C.TSC_TITLE} · {C.TSC_CODE} · {C.TSC_LEVEL}"),
            ("Mode of Delivery", "Classroom / Instructor-led (Physical)"),
            ("Course Duration",
             f"{C.TOTAL_HOURS} hours over {C.DAYS} days "
             f"({C.TRAINING_HOURS} h facilitation + {C.ASSESSMENT_HOURS} h assessment)"),
            ("Class Timing", f"Day 1: {_ampm(DAY1[0][0])} – {_ampm(DAY1[-1][1])}  ·  "
                             f"Day 2: {_ampm(DAY2[0][0])} – {_ampm(DAY2[-1][1])} "
                             f"(including the 2-hour assessment)"),
            ("Assessment Methods",
             "Written Assessment – Short Answer Questions (WA/SAQ), 1 hour · "
             "Case Study (CS), 1 hour · both open book"),
            ("Grading", C.ASSESSMENT["grading"]),
            ("Funding Validity", C.FUNDING_VALIDITY),
            ("Max Class Size", "20 learners")]
    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
    for k, v in info:
        cells = t.add_row().cells
        cells[0].text = ""; r = cells[0].paragraphs[0].add_run(k)
        r.bold = True; r.font.size = Pt(10); r.font.name = "Arial"
        P.shade_cell(cells[0], "F5F8FC")
        cells[1].text = ""; r2 = cells[1].paragraphs[0].add_run(v)
        r2.font.size = Pt(10); r2.font.name = "Arial"
        cells[0].width = Cm(6.0); cells[1].width = Cm(20.0)
    doc.add_paragraph()

    hdr(doc, "2. Learning Outcomes", 1)
    para(doc, "On completion of this course, learners will be able to:")
    for lo in C.LEARNING_OUTCOMES:
        p = doc.add_paragraph(style="List Bullet")
        P.md_runs(p, lo, size=10.5)

    hdr(doc, "3. Knowledge and Ability Statements Assessed", 1)
    t = doc.add_table(rows=0, cols=3); t.style = "Table Grid"
    hr = t.add_row().cells
    for i, h in enumerate(["Code", "Statement", "Assessed by"]):
        hr[i].text = ""; r = hr[i].paragraphs[0].add_run(h)
        r.bold = True; r.font.size = Pt(9.5); r.font.name = "Arial"
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        P.shade_cell(hr[i], "1F6FEB")
    for code, desc in C.TSC_KNOWLEDGE:
        cells = t.add_row().cells
        for i, v in enumerate([code, desc, "Written Assessment (SAQ)"]):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(v)
            r.font.size = Pt(9.5); r.font.name = "Arial"; r.bold = (i == 0)
    for code, desc in C.TSC_ABILITIES:
        cells = t.add_row().cells
        for i, v in enumerate([code, desc, "Case Study (CS)"]):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(v)
            r.font.size = Pt(9.5); r.font.name = "Arial"; r.bold = (i == 0)
    doc.add_paragraph()

    hdr(doc, "4. Course Outline", 1)
    for T in C.TOPICS:
        hdr(doc, f"Topic {T['num']}: {T['title']}", 2)
        para(doc, f"**{T['lu']}**  ·  Learning outcome **{T['lo']}**  ·  "
                  f"Knowledge/Ability: **{T['ka']}**", size=10.5)
        for st in T["subtopics"]:
            p = doc.add_paragraph(style="List Bullet")
            P.md_runs(p, st, size=10.5)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    hdr(doc, "5. Daily Schedule", 1)

    hdr(doc, "Day 1", 2)
    para(doc, f"**Total facilitation: {_f1 // 60} hours** ({_ampm(DAY1[0][0])} – "
              f"{_ampm(DAY1[-1][1])}, excluding a 45-minute lunch break and two 15-minute tea "
              f"breaks). Digital attendance is taken in the morning and the afternoon.",
         size=10)
    schedule_table(doc, DAY1)
    doc.add_paragraph()

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    hdr(doc, "Day 2", 2)
    para(doc, f"**Total facilitation: {_f2 // 60} hours + {_a2 // 60} hours assessment** "
              f"({_ampm(DAY2[0][0])} – {_ampm(DAY2[-1][1])}, excluding a 45-minute lunch break and "
              f"two 15-minute tea breaks). Assessment digital attendance is taken before "
              f"the assessment begins.", size=10)
    schedule_table(doc, DAY2)
    doc.add_paragraph()

    para(doc, f"**Total Training Hours: {C.TRAINING_HOURS} hours**   ·   "
              f"**Total Assessment Hours: {C.ASSESSMENT_HOURS} hours**   ·   "
              f"**Course Total: {C.TOTAL_HOURS} hours**", size=11)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    hdr(doc, "6. Hands-On Activities", 1)
    para(doc, "Each activity is a scenario-based case study run in small groups and "
              "debriefed by the trainer. Full step-by-step guidance, the scenario, the "
              "questions and the model answers are in the Learner Guide and in the "
              "standalone activity briefs in the activities/ folder.")
    from data_domain1 import DOMAIN1
    from data_domain2 import DOMAIN2
    from data_domain3 import DOMAIN3
    t = doc.add_table(rows=0, cols=5); t.style = "Table Grid"
    hr = t.add_row().cells
    for i, h in enumerate(["#", "Activity", "Topic", "Duration", "K/A assessed"]):
        hr[i].text = ""; r = hr[i].paragraphs[0].add_run(h)
        r.bold = True; r.font.size = Pt(9.5); r.font.name = "Arial"
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        P.shade_cell(hr[i], "1F6FEB")
    for a in DOMAIN1 + DOMAIN2 + DOMAIN3:
        cells = t.add_row().cells
        vals = [str(a["num"]), a["title"], f"Topic {int(a['topic'])}",
                f"{a['minutes']} min", a["ka"]]
        for i, v in enumerate(vals):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(v)
            r.font.size = Pt(9.5); r.font.name = "Arial"; r.bold = (i == 0)
    doc.add_paragraph()

    hdr(doc, "7. Assessment Plan", 1)
    para(doc, "**Written Assessment — Short Answer Questions (WA/SAQ)**: 4 questions, "
              "60 minutes, open book. Covers K1, K2, K3 and K4.", size=10.5)
    para(doc, "**Case Study (CS)**: 3 scenario-based questions, 60 minutes, open book. "
              "Covers A1, A2 and A3.", size=10.5)
    para(doc, "Both instruments are completed individually. Learners must answer all "
              "questions correctly to be assessed **Competent (C)**; otherwise **Not Yet "
              "Competent (NYC)**, with re-assessment offered in line with the Assessment "
              "Plan. An appeal process is available.", size=10.5)

    hdr(doc, "8. Resources Required", 1)
    for item in ["Trainer slide deck (PPT/PDF) and Learner Guide (PDF)",
                 "Activity briefs 1–3 (PDF) from the activities/ folder",
                 "Projector / TV, whiteboard and markers, flip chart for group work",
                 "Wi-Fi access to the LMS (lms-tms.tertiaryinfotech.com)",
                 "SSG digital attendance QR code (AM, PM and Assessment)",
                 "Assessment question papers, Assessment Plan and Assessment Summary Record",
                 "Copy of ISO 14001:2015 for trainer reference"]:
        p = doc.add_paragraph(style="List Bullet")
        P.md_runs(p, item, size=10.5)

    P.add_page_numbers(doc)
    P.enable_update_fields(doc)

    out = os.path.join(REPO, "courseware", f"LP-{C.SHORT_TITLE}.docx")
    doc.save(out)
    print("Saved", out)


if __name__ == "__main__":
    main()
