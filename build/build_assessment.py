#!/usr/bin/env python3
"""WSQ assessment set for TGS-2025054479 — ISO 14001:2015 Internal Auditor Training.

MIRRORS THE APPROVED ORIGINAL (pulled from the course's Drive Assessment folder; the
TMS record carries no assessment link, so the documented fallback applies):

  * Written Assessment (SAQ) — 4 questions, 60 minutes -> K1, K2, K3, K4
  * Case Study (CS)          — 3 questions, 60 minutes -> A1, A2, A3

Instrument types, question counts, timings and K/A mapping are UNCHANGED. Only the
content is rewritten from this course's slides, Learner Guide and activities.

Page layout (enforced):
  page 1 = cover naming the instrument
  page 2 = Trainee Information + Instructions to Candidate + Grading block, nothing else
  page 3+ = scenario and questions, each with a bordered answer box
There is NO "For Official Use Only" block at the back — the sign-off lives on page 2.
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import course_data as C
import prodoc as P

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, ".."))
OUT = os.path.join(REPO, "assessment")
ASSETS = os.path.join(REPO, "courseware", "assets")
LOGO = os.path.join(ASSETS, "tertiary-infotech-logo.png")
os.makedirs(OUT, exist_ok=True)

DARK = RGBColor(0x16, 0x1B, 0x26)
BRAND = RGBColor(0x1F, 0x6F, 0xEB)
GREY = RGBColor(0x55, 0x5B, 0x66)
LMS = "https://lms-tms.tertiaryinfotech.com/"


# ============================================================ doc helpers
def new_doc():
    d = Document()
    s = d.sections[0]
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2.0)
    n = d.styles["Normal"]
    n.font.name = "Arial"; n.font.size = Pt(11)
    return d


def line(d, text="", bold=False, size=11, color=DARK, after=6, align=None, italic=False):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = "Arial"
    return p


def hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url,
                          "http://schemas.openxmlformats.org/officeDocument/2006/"
                          "relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    nr = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    col = OxmlElement("w:color"); col.set(qn("w:val"), "1F6FEB")
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single")
    f = OxmlElement("w:rFonts"); f.set(qn("w:ascii"), "Arial"); f.set(qn("w:hAnsi"), "Arial")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "22")
    rPr.append(f); rPr.append(sz); rPr.append(col); rPr.append(u)
    nr.append(rPr)
    t = OxmlElement("w:t"); t.text = text; nr.append(t)
    h.append(nr); paragraph._p.append(h)


def cover(d, instrument, is_key=False):
    d.add_paragraph()
    if os.path.exists(LOGO):
        p = d.add_paragraph(); p.alignment = AL.CENTER
        p.add_run().add_picture(LOGO, width=Inches(2.1))
    line(d, C.ORG, bold=True, size=13, after=1, align=AL.CENTER)
    line(d, C.UEN, size=10, color=GREY, after=16, align=AL.CENTER)
    line(d, "ANSWER KEY" if is_key else "ASSESSMENT", bold=True, size=13,
         color=GREY, after=4, align=AL.CENTER)
    line(d, instrument, bold=True, size=24, color=BRAND, after=8, align=AL.CENTER)
    line(d, "For", size=12, color=GREY, after=8, align=AL.CENTER)
    line(d, C.TITLE, bold=True, size=18, after=4, align=AL.CENTER)
    line(d, f"TGS Ref No: {C.COURSE_CODE}", size=12, after=16, align=AL.CENTER)
    line(d, "Conducted by", size=11, color=GREY, after=2, align=AL.CENTER)
    line(d, C.ORG, bold=True, size=13, after=2, align=AL.CENTER)
    line(d, C.UEN, size=11, color=GREY, after=16, align=AL.CENTER)
    if is_key:
        line(d, "TRAINER / ASSESSOR COPY — NOT FOR CIRCULATION TO LEARNERS",
             bold=True, size=11, color=RGBColor(0xDC, 0x26, 0x26), after=6,
             align=AL.CENTER)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def admin_page(d, instrument, duration, nq, extra_instructions=()):
    """PAGE 2 — Trainee Information + Instructions + Grading. Nothing else."""
    line(d, "A: Trainee Information", bold=True, size=12, after=6)
    for t in ("Trainee Name (as per NRIC): ______________________________________",
              "Last three digits and alphabet of NRIC / FIN: ____________________",
              "Date: ______________________"):
        line(d, t, after=6)
    line(d, "", after=6)

    line(d, "B: Instructions to Candidate", bold=True, size=12, after=6)
    line(d, f"This is the {instrument}.", after=4)
    line(d, f"Duration: {duration}", bold=True, after=6)
    for i, t in enumerate([
        f"The assessor will pass the questions to you in hard copy. There are {nq} "
        f"questions. You must answer ALL questions.",
        "This is an OPEN BOOK assessment, to be completed INDIVIDUALLY. You may refer to "
        "the trainer slides, the Learner Guide, the activity briefs and your own notes.",
        "Write your answers in the answer box provided under each question. Use a black or "
        "blue pen; no liquid paper or correction tape.",
        "You need to answer all questions correctly to be assessed Competent (C).",
    ] + list(extra_instructions), 1):
        line(d, f"{i}. {t}", after=4)

    line(d, "", after=4)
    line(d, "Submission Procedure", bold=True, size=11, after=4)
    line(d, "1. Pass your completed hard copy to the assessor when time is up.", after=4)
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("2. Upload your answers to the LMS at ")
    r.font.size = Pt(11); r.font.name = "Arial"; r.font.color.rgb = DARK
    hyperlink(p, LMS, LMS)
    line(d, "", after=8)

    line(d, "C: Grading", bold=True, size=12, after=6)
    t = d.add_table(rows=0, cols=2); t.style = "Table Grid"
    for k, v in [("Result", "Competent (C)   /   Not Yet Competent (NYC)"),
                 ("Assessor Name", "_______________________________________"),
                 ("Assessor NRIC", "_______________________________________"),
                 ("Assessor Signature", "_______________________________________"),
                 ("Date", "_______________________________________"),
                 ("Assessor Feedback", " ")]:
        cells = t.add_row().cells
        cells[0].text = ""; rr = cells[0].paragraphs[0].add_run(k)
        rr.bold = True; rr.font.size = Pt(10.5); rr.font.name = "Arial"
        P.shade_cell(cells[0], "F5F8FC")
        cells[1].text = ""; r2 = cells[1].paragraphs[0].add_run(v)
        r2.font.size = Pt(10.5); r2.font.name = "Arial"
        cells[0].width = Cm(4.5); cells[1].width = Cm(12.0)
        if k == "Assessor Feedback":
            for _ in range(3):
                cells[1].add_paragraph()
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _cant_split(row):
    """Stop Word/LibreOffice splitting a table row across a page break."""
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit"); el.set(qn("w:val"), "true")
    trPr.append(el)


def _keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:keepNext"); el.set(qn("w:val"), "true")
    pPr.append(el)


def answer_box(d, lines=9, trailing=True):
    t = d.add_table(rows=1, cols=1); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cant_split(t.rows[0])
    c = t.rows[0].cells[0]
    c.text = ""
    p = c.paragraphs[0]
    r = p.add_run("Answer:")
    r.bold = True; r.font.size = Pt(10); r.font.name = "Arial"; r.font.color.rgb = GREY
    for _ in range(lines):
        c.add_paragraph()
    if trailing:
        d.add_paragraph()


def question(d, num, text, codes, scenario=None):
    if scenario:
        line(d, scenario, size=10.5, italic=True, after=8)
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f"Q{num}. ")
    r.bold = True; r.font.size = Pt(11.5); r.font.name = "Arial"
    r2 = p.add_run(text)
    r2.font.size = Pt(11.5); r2.font.name = "Arial"
    r3 = p.add_run(f"  ({codes})")
    r3.bold = True; r3.font.size = Pt(11.5); r3.font.name = "Arial"
    r3.font.color.rgb = BRAND
    _keep_with_next(p)


def key_answer(d, num, text, codes, ref):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"Q{num}  ({codes})")
    r.bold = True; r.font.size = Pt(12); r.font.name = "Arial"; r.font.color.rgb = BRAND
    line(d, f"Taught in: {ref}", size=9.5, color=GREY, italic=True, after=6)
    t = d.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.rows[0].cells[0]; c.text = ""
    P.shade_cell(c, "F5F8FC")
    first = c.paragraphs[0]
    rr = first.add_run("Suggested answer (not exhaustive) — award Competent where the "
                       "learner covers the substance of the points below:")
    rr.bold = True; rr.font.size = Pt(10); rr.font.name = "Arial"
    for blk in text.split("\n\n"):
        if not blk.strip():
            continue
        pp = c.add_paragraph()
        pp.paragraph_format.space_after = Pt(5)
        P.md_runs(pp, blk.strip(), size=10.5)
    d.add_paragraph()


def footer(d, label):
    sec = d.sections[0]
    f = sec.footer; f.is_linked_to_previous = False
    p = f.paragraphs[0]; p.alignment = AL.CENTER; p.text = ""
    r = p.add_run(f"{C.TITLE}  ·  {C.COURSE_CODE}  ·  {label}     Page ")
    r.font.size = Pt(8); r.font.color.rgb = GREY; r.font.name = "Arial"
    P._field(p, "PAGE", "1").font.size = Pt(8)
    r2 = p.add_run(" of "); r2.font.size = Pt(8); r2.font.color.rgb = GREY
    r2.font.name = "Arial"
    P._field(p, "NUMPAGES", "1").font.size = Pt(8)


# ============================================================ CONTENT — WA (SAQ)
WA_DURATION = "60 minutes"
WRITTEN = [
    dict(
        n=1, codes="K1",
        scenario="You have recently joined a mid-sized manufacturing company as an internal "
                 "auditor. Your audit manager has asked you to familiarise yourself with the "
                 "International Professional Practices Framework (IPPF) before conducting "
                 "your first ISO 14001:2015 environmental management system audit. During a "
                 "team briefing, a colleague asks you to explain the framework.",
        q="Describe the key components of the International Professional Practices "
          "Framework (IPPF), including the definition of Internal Auditing, the Code of "
          "Ethics, the core principles and the International Standards for the Professional "
          "Practice of Internal Auditing. Explain how each component guides internal "
          "auditors in their professional practice.",
        ref="Topic 1, Learner Guide §2.1; slides — 'The IPPF — Mandatory Guidance', "
            "'The Code of Ethics', 'Independence and Objectivity'",
        a="**The IPPF comprises four mandatory components** that together define what an "
          "internal auditor is, how they must behave, and how they must work.\n\n"
          "**1. The Definition of Internal Auditing.** Internal auditing is an "
          "*independent, objective assurance and consulting activity designed to add value "
          "and improve an organisation's operations*. It helps the organisation accomplish "
          "its objectives by bringing a systematic, disciplined approach to evaluating and "
          "improving the effectiveness of risk management, control and governance "
          "processes. **How it guides practice:** it establishes that the auditor is not an "
          "inspector or a policeman — the purpose is to add value, and the work must be "
          "both independent and objective.\n\n"
          "**2. The Code of Ethics.** Four principles and their rules of conduct: "
          "**Integrity** (perform work with honesty, diligence and responsibility; do not "
          "knowingly be party to illegal activity); **Objectivity** (maintain an unbiased "
          "mental attitude, avoid conflicts of interest, and do not audit an area for which "
          "you hold operational responsibility); **Confidentiality** (respect the value and "
          "ownership of information received and do not use it for personal gain); and "
          "**Competency** (apply the knowledge, skills and experience required, and decline "
          "or obtain support for work beyond your competence). **How it guides practice:** "
          "it determines who may audit what — for example, an auditor may not audit their "
          "own department, and must bring in a specialist where the process is technically "
          "beyond them.\n\n"
          "**3. The Core Principles.** These articulate the characteristics of an effective "
          "internal audit function: demonstrating integrity, competence and due "
          "professional care; being objective and free from undue influence; aligning with "
          "the strategies, objectives and risks of the organisation; being appropriately "
          "positioned and adequately resourced; communicating effectively; providing "
          "risk-based assurance; and being insightful, proactive and future-focused. **How "
          "it guides practice:** the audit plan must follow organisational risk rather than "
          "organisational convenience.\n\n"
          "**4. The International Standards for the Professional Practice of Internal "
          "Auditing.** Attribute standards describe the characteristics of the function and "
          "the auditors; performance standards describe the nature of the work and the "
          "criteria against which it is measured. **How it guides practice:** engagements "
          "must be planned, evidenced, documented and reported to a consistent professional "
          "standard.\n\n"
          "**Applied to an ISO 14001:2015 EMS audit**, the IPPF supplies the professional "
          "conduct framework while ISO 14001 clause 9.2 supplies the requirement to audit "
          "and ISO 19011:2018 supplies the method. Credit answers that connect the "
          "components to concrete auditor behaviour, not merely list them."),
    dict(
        n=2, codes="K2",
        scenario="A chemical processing plant is preparing for an internal audit of its "
                 "Environmental Management System (EMS) aligned with ISO 14001:2015. As the "
                 "assigned internal auditor you must understand the organisation's context "
                 "before developing the audit plan. The plant manager has given you the "
                 "strategic plan, the organisation chart, operational process maps and the "
                 "register of environmental aspects and impacts.",
        q="Explain why it is essential for an internal auditor to understand an "
          "organisation's significant risks, objectives, resources, operations, businesses "
          "and structure before conducting an ISO 14001:2015 EMS audit. Provide specific "
          "examples of how each element informs the audit planning process.",
        ref="Topic 1, Learner Guide §2.3–2.5; slides — 'Clause 4 — Context of the "
            "Organisation', 'Clause 6.1.2 — Environmental Aspects', 'Risks and "
            "Opportunities are NOT Aspects'",
        a="Understanding the organisation is what converts a generic clause checklist into "
          "a **risk-based audit**. ISO 14001:2015 clause 4.1 requires the organisation to "
          "determine its internal and external issues, and clause 9.2.2 requires the audit "
          "programme to take into account **the environmental importance of the processes "
          "concerned** and **changes affecting the organisation**. An auditor who does not "
          "understand the organisation cannot satisfy either.\n\n"
          "**Significant risks.** The EMS exists to control significant environmental "
          "aspects and to fulfil compliance obligations. Knowing that the plant's dominant "
          "risks are chemical spills, air emissions and hazardous waste disposal tells the "
          "auditor where to spend audit time and which controls (clause 8.1) and monitoring "
          "(clause 9.1) must be verified. *Example:* a bunded solvent tank farm justifies "
          "far more audit effort than the administration office.\n\n"
          "**Objectives.** Environmental objectives (clause 6.2) must be consistent with "
          "the policy and take account of significant aspects and compliance obligations. "
          "Knowing the objectives lets the auditor test whether they are measurable, "
          "monitored and resourced, and whether they trace back to a significant aspect. "
          "*Example:* a 20% carbon reduction target within two years lets the auditor ask "
          "for the indicator, the baseline, the owner and the evidence of progress under "
          "clause 6.2.2.\n\n"
          "**Resources.** Clause 7.1 requires the organisation to determine and provide the "
          "resources the EMS needs, and clause 7.2 requires competence. *Example:* the "
          "auditor checks whether a competent environmental officer is appointed and "
          "whether budget exists for pollution control equipment — a vacant EHS role is "
          "itself an audit lead.\n\n"
          "**Operations.** Knowing the processes tells the auditor which activities create "
          "the most significant aspects, and lets them follow a real audit trail. "
          "*Example:* a process using volatile solvents directs the auditor to emission "
          "controls, operating criteria and the applicable air-quality obligations.\n\n"
          "**Business and structure.** The organisation chart establishes who holds "
          "authority and responsibility (clause 5.3), who to interview, and where "
          "independence problems could arise. *Example:* the auditor confirms that "
          "responsibility for reporting EMS performance to top management has been assigned "
          "and communicated, and confirms they are not auditing their own function.\n\n"
          "Credit answers that link each element to a **specific clause and a specific "
          "planning decision**, rather than describing the organisation generically."),
    dict(
        n=3, codes="K3",
        scenario="You are an internal auditor at a logistics company that has implemented "
                 "ISO 14001:2015 across its warehousing and transportation divisions. Your "
                 "audit lead has asked you to review the company's internal audit "
                 "methodology before the annual EMS audit cycle, to ensure it is robust and "
                 "consistent with both ISO 14001:2015 and the organisation's own procedures.",
        q="Describe the key internal audit methodologies and procedures that an "
          "organisation should have in place for conducting ISO 14001:2015 EMS audits. "
          "Explain how these methodologies ensure a systematic and effective audit process.",
        ref="Topic 2, Learner Guide §3.1–3.5; slides — 'The Seven Audit Principles', "
            "'Audit Programme vs Individual Audit', 'The Six-Stage Audit Lifecycle', "
            "'Building the Audit Checklist', 'Audit Evidence'",
        a="**1. A documented internal audit procedure.** Defining scope, frequency, "
          "methods, responsibilities, planning requirements and reporting — required by "
          "ISO 14001:2015 clause 9.2.2 and guided by ISO 19011:2018. It must be kept "
          "current; a procedure that predates a major change in the organisation is itself "
          "a finding.\n\n"
          "**2. An audit programme, distinct from the individual audit.** ISO 19011 "
          "Clause 5 governs the programme — establishing objectives, determining programme "
          "risks and opportunities, resourcing it, implementing, monitoring, reviewing and "
          "improving it. Clause 9.2.2 requires the programme to consider the environmental "
          "importance of the processes, changes affecting the organisation, and the results "
          "of previous audits. **Effect:** audits are scheduled by risk rather than "
          "allocated equally, and the programme improves year on year.\n\n"
          "**3. Risk-based audit planning.** High-risk areas — significant aspects, "
          "licensed discharges, recent process changes, areas with previous findings — are "
          "audited more often and more deeply than low-risk administrative processes.\n\n"
          "**4. The six-stage audit lifecycle (ISO 19011 §6.2–6.7).** Initiating (contact, "
          "authority, feasibility); preparing (document review, audit plan, work "
          "assignment, checklist, sampling plan); conducting (opening meeting, collecting "
          "and **verifying** information, generating findings, closing meeting); preparing "
          "and distributing the report; completing the audit; and conducting **follow-up**. "
          "**Effect:** nothing is improvised, and every audit is repeatable.\n\n"
          "**5. Work documents — the checklist.** Organisation-specific and risk-weighted, "
          "with the clause, the question and the **evidence expected**. A generic "
          "clause-by-clause tick-list produces the same findings every year and misses what "
          "actually changed.\n\n"
          "**6. Evidence and sampling rules.** Objective evidence from four sources — "
          "interview, observation, documents and records — that is sufficient, appropriate "
          "and traceable. The sampling basis must be recorded and stated in the report, "
          "because an audit examines only a sample.\n\n"
          "**7. Auditor competence and impartiality.** Clause 9.2.2 requires auditors to be "
          "selected so as to ensure objectivity and the impartiality of the process; "
          "auditors must not audit their own work.\n\n"
          "**8. Reporting, corrective action and follow-up.** Findings graded and reported "
          "to relevant management; corrective action under clause 10.2 addressing the "
          "**cause**, with the **effectiveness reviewed** under 10.2(d) and the results "
          "retained. **Effect:** closure means verified, not merely marked closed.\n\n"
          "Credit answers that distinguish the **programme** from the **individual audit** "
          "and that name follow-up/effectiveness verification."),
    dict(
        n=4, codes="K4",
        scenario="A food manufacturing company is expanding operations to include a new "
                 "production line that uses refrigerants and generates significant "
                 "wastewater. The EMS manager has asked the internal audit team to conduct "
                 "a risk assessment to identify the environmental risks associated with the "
                 "expansion. You must apply appropriate risk identification and assessment "
                 "techniques as part of the preliminary audit survey.",
        q="Explain the key risk identification and assessment techniques used in the "
          "context of ISO 14001:2015 internal auditing. Describe how these techniques help "
          "auditors evaluate environmental risks and prioritise audit focus areas.",
        ref="Topic 3, Learner Guide §4.1; slides — 'Evaluating Significance — the 5×5 "
            "Matrix', 'The Override Rules', 'Three Significance Methods', 'FMEA for "
            "Environmental Risk'",
        a="**1. Environmental aspect and impact analysis (clause 6.1.2).** Identify the "
          "activities, products and services that interact with the environment, determine "
          "the associated impacts, and evaluate significance **using established "
          "criteria**. The aspect is the cause the organisation controls (*discharge of "
          "effluent containing copper*); the impact is the effect on the environment "
          "(*watercourse contamination*). A complete register covers **normal, abnormal and "
          "reasonably foreseeable emergency** conditions, past/present/planned activities, "
          "and both direct and indirect aspects.\n\n"
          "**2. Identification methods.** Process mapping; **input–output (mass balance) "
          "analysis** — unaccounted input usually reveals an unregistered fugitive "
          "emission; site walkthrough of drains, stacks, bunds and storage; and life cycle "
          "stages to catch indirect aspects.\n\n"
          "**3. Risk matrices (likelihood × consequence).** A 5×5 matrix scores severity "
          "against likelihood, banding results Low / Medium / High / Extreme with a defined "
          "significance threshold (commonly ≥ 10). *Example:* a refrigerant leak may be of "
          "moderate likelihood but high consequence, placing it in the significant band. "
          "The matrix makes priorities visible and communicable.\n\n"
          "**4. Override rules.** An aspect is significant **regardless of its score** if a "
          "legal limit or licence condition applies, a regulator has acted within the last "
          "three years, a substantiated complaint exists, or a controlled hazardous "
          "substance or toxic industrial waste is involved. Without these, a licensed "
          "discharge can score 'medium' and wrongly fall off the significant list.\n\n"
          "**5. Multi-criteria scoring.** Severity × Frequency × Legal exposure × "
          "Stakeholder concern builds compliance obligations (6.1.3) and interested parties "
          "(4.2) explicitly into the evaluation — suited to larger organisations.\n\n"
          "**6. FMEA (Severity × Occurrence × Detection = RPN).** Adds the dimension the "
          "others miss — whether the failure can be **detected before** the impact occurs. "
          "*Example:* a cracked bund with S=9, O=4, D=7 gives RPN 252; adding leak "
          "detection reduces D to 2 and RPN to 72. Improving detection is usually the "
          "cheapest risk reduction, and it links directly to clause 9.1 monitoring.\n\n"
          "**7. Supporting techniques.** SWOT for context; process mapping and flowcharts; "
          "review of the legal register, incident and complaint records; and the previous "
          "audit results.\n\n"
          "**How this prioritises audit focus.** Aspects that are significant, recently "
          "changed, legally constrained or previously non-conforming receive the most audit "
          "time. Credit answers that state a **threshold** and explain that ISO 14001 "
          "prescribes no method — only that the criteria are established, documented and "
          "consistently applied."),
]

# ============================================================ CONTENT — Case Study
CS_DURATION = "60 minutes"

CS_ITEMS = [
    dict(
        n=1, codes="A1",
        scenario=(
            "**Scenario 1 — GreenPack Manufacturing Pte Ltd**\n\n"
            "GreenPack Manufacturing Pte Ltd is a Singapore-based company producing "
            "biodegradable packaging materials. The company achieved ISO 14001:2015 "
            "certification two years ago and employs 150 staff across three departments: "
            "Production, Warehousing & Logistics, and Administration.\n\n"
            "Recently the company expanded its production capacity by 40% and introduced a "
            "new thermal lamination process that uses solvent-based adhesives. The Managing "
            "Director has noticed that energy consumption has increased by 25%, volatile "
            "organic compound (VOC) emissions have risen, and three customer complaints "
            "about chemical odours from the packaging have been received.\n\n"
            "The company's Environmental Policy was last reviewed 18 months ago and does "
            "not reference the new process. The environmental aspects and impacts register "
            "has not been updated since the expansion. The internal audit team has been "
            "directed to conduct a comprehensive EMS audit.\n\n"
            "As the lead internal auditor you have access to the Environmental Policy, the "
            "current aspects and impacts register, the organisational chart showing "
            "reporting lines, and the strategic plan, which includes an objective to reduce "
            "the carbon footprint by 30% within three years. The company follows the "
            "International Professional Practices Framework (IPPF) for its internal audit "
            "function."),
        q="Based on the GreenPack Manufacturing scenario, prepare a report that addresses "
          "the following:\n\n"
          "(a) Identify which ISO 14001:2015 clauses and IPPF standards are most relevant "
          "to this audit engagement, and explain why each is applicable to GreenPack's "
          "current situation.\n\n"
          "(b) Analyse how GreenPack's organisational risks, objectives, resources and "
          "operational changes should influence the scope and focus of the internal audit "
          "plan.",
        ref="Activity 1 (GreenPack) · Learner Guide §2.1–2.5 and §5, Activity 1",
        a="**(a) Relevant clauses and IPPF standards.**\n\n"
          "**IPPF.** The **Definition of Internal Auditing** frames the engagement as "
          "independent, objective assurance that adds value. The **Code of Ethics** binds "
          "the auditor to Integrity, **Objectivity** (the auditor must not have operational "
          "responsibility for Production), Confidentiality and **Competency** (solvent "
          "emissions may require technical support). The **Core Principles** and "
          "**International Standards** require the engagement to be risk-based, adequately "
          "planned, evidenced and reported.\n\n"
          "**ISO 14001:2015 clauses, each tied to a scenario fact:**\n\n"
          "• **6.1.2 Environmental aspects — the strongest finding.** The clause explicitly "
          "requires aspects to be determined considering **change, including new or "
          "modified activities**, and abnormal and reasonably foreseeable emergency "
          "conditions. The register has not been updated since a 40% expansion and a new "
          "solvent process. Likely **major**.\n\n"
          "• **5.2 Environmental policy.** Not reviewed for 18 months and silent on the new "
          "process, despite the requirement that the policy be appropriate to the purpose "
          "and context of the organisation, including the nature, scale and environmental "
          "impacts of its activities.\n\n"
          "• **4.1 Context.** The new process, rising energy cost and capacity change are "
          "internal issues; air-quality regulation and the landlord's lease are external.\n\n"
          "• **4.2 Interested parties.** Customers (three complaints), the neighbouring "
          "tenant, the regulator and the landlord — and clause 4.2(c) determines which of "
          "their expectations become compliance obligations.\n\n"
          "• **4.3 Scope.** A 40% capacity increase and a new process test whether the "
          "declared EMS scope is still valid.\n\n"
          "• **6.1.3 Compliance obligations.** VOC emissions engage air-quality "
          "obligations, which must be identified and their applicability determined.\n\n"
          "• **6.2 Objectives.** A 30% carbon-reduction objective is contradicted by a 25% "
          "rise in energy consumption — test the indicator, the resources and the "
          "monitoring.\n\n"
          "• **7.4 Communication** (handling of the complaints), **8.1 Operational "
          "control** (operating criteria for the lamination line), **9.1.1 Monitoring** "
          "(is VOC and energy performance measured at all?) and **9.3 Management review** "
          "(did it consider the change?).\n\n"
          "**(b) How the organisation shapes the scope.**\n\n"
          "**Objective:** determine whether the EMS conforms to ISO 14001:2015 and to "
          "GreenPack's own requirements, and is effectively implemented and maintained, "
          "with particular regard to the capacity expansion and the lamination process. "
          "**Criteria:** ISO 14001:2015, GreenPack's own EMS procedures, and its applicable "
          "compliance obligations. **Boundaries:** the site, all three departments, EMS "
          "activity from the date of the expansion.\n\n"
          "**Risks** drive the priority order: (1) the lamination line — new, generating "
          "complaints, legally exposed; (2) the aspect-identification and change-management "
          "process itself, which is the systemic cause; (3) policy, objectives and "
          "management review; (4) energy performance against the carbon objective; "
          "(5) warehousing; (6) administration. **Objectives** direct the auditor to the "
          "6.2.2 evidence (what, resources, who, when, how evaluated). **Resources** raise "
          "clause 7.1 and 7.2 questions about competence for the new process. "
          "**Operational change** is the thread connecting everything — it should have "
          "triggered a re-evaluation of aspects, controls, obligations and the policy.\n\n"
          "**Allocating equal time to all three departments would be indefensible**: audit "
          "effort must follow environmental risk and change. Credit answers that state this "
          "explicitly and that record the auditor's independence."),
    dict(
        n=2, codes="A2",
        scenario=(
            "**Scenario 2 — AquaPure Water Treatment Services Ltd**\n\n"
            "AquaPure Water Treatment Services Ltd provides industrial water treatment "
            "solutions to factories and commercial buildings across Southeast Asia. The "
            "company operates a central treatment plant and deploys mobile treatment units "
            "to client sites. AquaPure recently won a large government contract to treat "
            "wastewater from a new industrial park, which has tripled its operational "
            "workload. The EMS Manager has scheduled an internal audit of the wastewater "
            "treatment operations.\n\n"
            "As the assigned internal auditor you discover the following during your "
            "preliminary review:\n\n"
            "• The company's internal audit procedure was written five years ago and has "
            "not been updated to reflect the expanded operations or the government contract "
            "requirements.\n"
            "• The previous audit report (12 months ago) identified two minor "
            "nonconformities: (1) incomplete training records for mobile unit operators, "
            "and (2) inconsistent calibration of pH monitoring equipment. The corrective "
            "actions were marked as 'closed' but no verification evidence was found.\n"
            "• AquaPure engages four key stakeholder groups in its audit process: "
            "Operations staff, Client representatives, Regulatory authorities (e.g. NEA), "
            "and Senior management.\n"
            "• The audit schedule allocates equal time to all departments, regardless of "
            "risk level or operational significance.\n"
            "• The Operations Manager has expressed concern that mobile unit operators at "
            "client sites are not consistently following the effluent discharge monitoring "
            "procedures."),
        q="Based on the AquaPure scenario, analyse the following:\n\n"
          "(a) Identify and explain the key stages of an ISO 14001:2015 internal audit "
          "engagement that should be followed for AquaPure's upcoming audit. For each "
          "stage, describe specific actions the auditor should take relevant to AquaPure's "
          "situation.\n\n"
          "(b) Evaluate how the auditor should engage with each of the four identified "
          "stakeholder groups throughout the audit process to ensure a comprehensive and "
          "effective audit outcome.",
        ref="Activity 2 (AquaPure) · Learner Guide §3.2–3.8 and §5, Activity 2",
        a="**(a) The stages, applied to AquaPure.**\n\n"
          "**Stage 1 — Initiating (ISO 19011 §6.2).** Establish contact with the EMS "
          "Manager and confirm authority to audit. **Confirm feasibility**: the mobile "
          "units operate on **client premises**, so site access must be negotiated in "
          "advance. If access cannot be obtained, the audit is not feasible as scoped and "
          "this must be raised now, not in the report.\n\n"
          "**Stage 2 — Preparing (§6.3).** Review the five-year-old audit procedure, the "
          "previous audit report, both CAR records, the discharge licence and the "
          "monitoring data. Prepare the audit plan (objective, criteria, scope, locations "
          "including a sample of client sites, dates, roles, opening and closing meetings). "
          "Prepare a **risk-weighted checklist** targeting mobile operations, calibration, "
          "training and the two previous nonconformities. **Define the sample** — for "
          "example three of the mobile units across two client sites, plus six months of pH "
          "and discharge records — and be ready to justify it.\n\n"
          "**Stage 3 — Conducting (§6.4).** Hold the opening meeting (confirm plan, "
          "criteria, sampling basis and its uncertainty, grading, escort and safety). "
          "Collect evidence by interview, observation, documents and records — critically, "
          "**observe a mobile unit actually performing discharge monitoring** rather than "
          "relying on what operators say. Verify before treating information as evidence. "
          "Generate findings against the criteria and hold the closing meeting.\n\n"
          "**Stage 4 — Reporting (§6.5).** Report the criteria, scope, **the sampling "
          "statement**, the findings each with evidence and grade, and the conclusion on "
          "conformity **and effective implementation**.\n\n"
          "**Stage 5 — Completing (§6.6).** Distribute, retain the records, close the "
          "audit.\n\n"
          "**Stage 6 — Follow-up (§6.7).** Verify corrective actions are implemented "
          "**and effective** — the stage AquaPure evidently skipped last year.\n\n"
          "**Programme-level findings (ISO 14001 cl. 9.2.2).** The audit procedure has not "
          "been maintained despite the workload tripling; equal time regardless of risk "
          "contradicts the requirement to consider the environmental importance of the "
          "processes; and closing both prior NCs without verification evidence breaches "
          "clause **10.2(d)** and its retention requirement. Strong answers separate these "
          "**programme** failures from findings about this individual audit.\n\n"
          "**(b) Engaging the four stakeholder groups.**\n\n"
          "**Operations staff** — the primary evidence source and the people who know what "
          "actually happens. Interview at their normal workplace, put them at ease, explain "
          "why you are there and why you take notes, use open questions, and never "
          "apportion blame. They will reveal the gap between the procedure and the practice "
          "— exactly the concern the Operations Manager raised.\n\n"
          "**Client representatives** — **not the auditee**. They are gatekeepers for "
          "access to the mobile units. Obtain permission in advance, comply with their site "
          "rules and induction, and protect confidentiality in both directions. The auditor "
          "has no authority over them and cannot raise findings against them.\n\n"
          "**Regulatory authorities (NEA)** — engaged **indirectly**. Review licences, "
          "permit conditions, submitted monitoring returns, correspondence and any notices. "
          "An internal auditor does **not** interview the regulator; the regulator's "
          "records are the evidence.\n\n"
          "**Senior management** — engaged at the opening and closing meetings, and as the "
          "audience for the programme-level findings. They own the resources (clause 7.1), "
          "the leadership requirements (clause 5.1) and the decision to move to risk-based "
          "scheduling. Report to them the systemic issues, not just the individual "
          "nonconformities.\n\n"
          "Credit answers that **differentiate** the four groups by their role, timing and "
          "the risk of engaging them badly."),
    dict(
        n=3, codes="A3",
        scenario=(
            "**Scenario 3 — EcoTech Electronics Pte Ltd**\n\n"
            "EcoTech Electronics Pte Ltd is a contract electronics manufacturer that "
            "assembles printed circuit boards (PCBs) and electronic components for "
            "international clients. The company operates a 10,000 sqm facility in Tuas, "
            "Singapore, with 200 employees. EcoTech's operations generate several "
            "environmental concerns including hazardous waste (solder dross, flux residues "
            "and waste solvents), wastewater containing heavy metals (lead, tin, copper) "
            "from the PCB washing process, and air emissions from soldering operations.\n\n"
            "EcoTech recently failed a regulatory inspection by the National Environment "
            "Agency (NEA), receiving a warning for exceeding the permitted trade effluent "
            "discharge limits for copper concentration. Additionally, a routine check by "
            "the facilities team found that three hazardous waste storage containers in the "
            "soldering area were unlabelled, and the spill containment bund was cracked and "
            "leaking.\n\n"
            "The company's risk register was last updated eight months ago and does not "
            "reflect recent changes including: (1) the introduction of a lead-free "
            "soldering process that uses higher temperatures and different flux chemicals, "
            "(2) a 30% increase in production volume due to a new client contract, and "
            "(3) the departure of the Environmental Health & Safety (EHS) Manager three "
            "months ago, whose position remains vacant.\n\n"
            "As part of the upcoming ISO 14001:2015 internal audit you have been tasked "
            "with conducting a preliminary risk assessment of EcoTech's environmental "
            "management practices and preparing an audit report with corrective action "
            "recommendations."),
        q="Based on the EcoTech Electronics scenario, complete the following:\n\n"
          "(a) Conduct a preliminary risk assessment by identifying at least five "
          "significant environmental risks facing EcoTech, assessing each risk using a "
          "likelihood and consequence framework, and prioritising them for audit "
          "attention.\n\n"
          "(b) For the three most critical risks identified, write effective corrective "
          "action requests (CARs) that include the nonconformity description, the relevant "
          "ISO 14001:2015 clause, the root cause analysis, and specific corrective and "
          "preventive actions.\n\n"
          "(c) Outline the key sections and content that should be included in the audit "
          "report for EcoTech's management.",
        ref="Activity 3 (EcoTech) · Learner Guide §4.1–4.4 and §5, Activity 3",
        a="**(a) Preliminary risk assessment** (5×5 matrix, score = severity × likelihood; "
          "significant at ≥ 10; override rules apply).\n\n"
          "**1. Trade effluent copper exceedance.** Aspect: discharge of trade effluent "
          "containing copper from PCB washing. Impact: watercourse contamination and "
          "aquatic toxicity. **S=4, L=4 → 16 (Extreme).** Automatically significant: a "
          "legal limit was breached **and** a regulator issued a warning. Clauses 6.1.3, "
          "8.1, 9.1.2.\n\n"
          "**2. Cracked, leaking spill containment bund.** Aspect: potential release of "
          "solvents and metal-bearing liquids — an **emergency** condition. Impact: soil "
          "and groundwater contamination. **S=5, L=3 → 15 (Extreme).** The designed control "
          "has already failed. Clauses 8.1, 8.2.\n\n"
          "**3. Three unlabelled hazardous waste containers.** Aspect: storage of toxic "
          "industrial waste without identification. Impact: mis-handling, incompatible "
          "mixing, improper disposal. **S=4, L=4 → 16 (Extreme).** Automatically "
          "significant — toxic industrial waste. Clauses 8.1, 7.2, 7.3.\n\n"
          "**4. Risk/aspects register eight months stale across three material changes.** "
          "The lead-free process changes flux chemistry and the emission profile; +30% "
          "volume changes every loading rate. **S=4, L=5 → 20 (Extreme).** Clause 6.1.2 "
          "explicitly requires change to be considered.\n\n"
          "**5. EHS Manager post vacant for three months.** No owner for aspects, the legal "
          "register, monitoring or compliance evaluation. **S=3, L=5 → 15 (Extreme).** "
          "Clauses 5.3, 7.1, 7.2. This is the **resource gap underlying risks 1–4**.\n\n"
          "Priority for audit attention: 4 and 1 first (systemic and legally breached), "
          "then 2 and 3 (failed operational controls), then 5 as the underlying cause. "
          "Credit answers that separate aspect from impact correctly, cover normal/abnormal/"
          "emergency conditions, and apply the **override rules** so the copper exceedance "
          "is significant regardless of arithmetic.\n\n"
          "**(b) Three corrective action requests.**\n\n"
          "**CAR 1 — MAJOR.** *Requirement:* ISO 14001:2015 cl. 6.1.3 requires compliance "
          "obligations to be determined and taken into account, and cl. 8.1 requires "
          "operational control against established operating criteria; the trade effluent "
          "discharge is subject to a permitted copper limit. *Evidence:* NEA inspection "
          "warning recording copper above the permitted limit at the site discharge point; "
          "pre-treatment plant records show no capacity re-assessment after the 30% "
          "production increase. *Nonconformity:* operational control is not sufficient to "
          "fulfil the compliance obligation for trade effluent copper. *Root cause "
          "(5 Whys):* exceedance → pre-treatment not removing copper → loading rose after "
          "the volume increase → capacity never re-assessed → **the management-of-change "
          "procedure contains no step requiring environmental re-assessment**. *Correction:* "
          "re-dose and re-test the effluent; notify NEA of the actions taken. *Corrective "
          "action:* amend the management-of-change procedure to mandate environmental and "
          "capacity re-assessment for any volume or process change; re-assess pre-treatment "
          "capacity at the new loading; install continuous monitoring where practicable. "
          "*Owner and date:* named role, real date. *Effectiveness verification:* review "
          "three months of discharge results against the limit.\n\n"
          "**CAR 2 — MAJOR.** *Requirement:* cl. 6.1.2 requires aspects to be determined "
          "considering change, including new or modified activities, and reasonably "
          "foreseeable emergency situations. *Evidence:* the aspects and risk register is "
          "dated eight months prior and contains no entry for the lead-free soldering "
          "process, the 30% volume increase or the revised flux chemistry. *Nonconformity:* "
          "environmental aspects have not been re-determined following material changes. "
          "*Root cause:* the same change-management gap, compounded by the vacant EHS "
          "Manager role leaving the register unowned. *Correction:* re-run aspect "
          "identification for all three changes. *Corrective action:* assign ownership; make "
          "aspects review a mandatory gate in the change process; set a review frequency. "
          "*Effectiveness verification:* confirm at the next change that the gate fired.\n\n"
          "**CAR 3 — MINOR, escalating to MAJOR.** *Requirement:* cl. 8.1 requires "
          "operational controls consistent with operating criteria; hazardous waste storage "
          "must permit correct identification and handling. *Evidence:* three waste "
          "containers in the soldering area observed without labels identifying contents or "
          "hazards; the adjacent spill containment bund cracked with visible leakage. "
          "*Nonconformity:* operational controls for hazardous waste storage and spill "
          "containment are not implemented as required. *Root cause:* no scheduled "
          "inspection of the waste storage area and no bund integrity testing regime; "
          "competence and awareness gaps after the EHS Manager's departure. *Correction:* "
          "label the containers; repair the bund; contain and clean the leak. *Corrective "
          "action:* documented weekly waste-area inspection; bund integrity testing added "
          "to the maintenance schedule; re-train soldering staff on waste labelling and "
          "**evaluate the effectiveness of that training** (cl. 7.2). *Grade rationale:* "
          "minor individually, but raised to **major** because it occurs together with "
          "CAR 1 and CAR 2, evidencing a systemic operational-control breakdown.\n\n"
          "Reject root causes that stop at 'the operator was careless' — a root cause must "
          "be a **process or system** gap. Reject CARs that list only corrections.\n\n"
          "**(c) The audit report — key sections.**\n\n"
          "1. Audit objectives, scope and criteria, including anything excluded. "
          "2. Audit team, auditees and dates. 3. **The sampling basis and its "
          "limitations.** 4. The findings, each with its objective evidence, the clause and "
          "its grade. 5. Positive observations. 6. Opportunities for improvement, kept "
          "**separate** from nonconformities. 7. **The audit conclusion** — on conformity "
          "**and** on whether the EMS is effectively implemented and maintained "
          "(cl. 9.2.1 requires both). 8. Agreed corrective actions, owners and due dates. "
          "9. Distribution list and follow-up arrangements.\n\n"
          "Credit answers that include the **sampling statement** — it is the most commonly "
          "omitted element and it protects both the auditor and the organisation."),
]


# ============================================================ coverage guard
def check_coverage():
    ks = {c for c, _ in C.TSC_KNOWLEDGE}
    as_ = {c for c, _ in C.TSC_ABILITIES}
    wa = set()
    for q in WRITTEN:
        wa |= {x.strip() for x in q["codes"].split(",")}
    cs = set()
    for q in CS_ITEMS:
        cs |= {x.strip() for x in q["codes"].split(",")}
    print("  WA covers:", ", ".join(sorted(wa)), "  (required:", ", ".join(sorted(ks)), ")")
    print("  CS covers:", ", ".join(sorted(cs)), "  (required:", ", ".join(sorted(as_)), ")")
    missing_k = ks - wa
    missing_a = as_ - cs
    if missing_k or missing_a:
        raise SystemExit(f"COVERAGE FAILURE — missing K: {sorted(missing_k)} "
                         f"missing A: {sorted(missing_a)}")
    extra = (wa - ks) | (cs - as_)
    if extra:
        raise SystemExit(f"COVERAGE FAILURE — unknown codes used: {sorted(extra)}")
    print("  coverage OK — every K in the WA, every A in the CS")


# ============================================================ builders
def build_wa_paper():
    d = new_doc()
    cover(d, "Written Assessment (SAQ)")
    admin_page(d, "Written Assessment — Short Answer Questions (SAQ)",
               WA_DURATION, len(WRITTEN))
    line(d, "C: Questions", bold=True, size=12, after=10)
    for i, q in enumerate(WRITTEN):
        question(d, q["n"], q["q"], q["codes"], scenario=q["scenario"])
        answer_box(d, lines=10, trailing=(i < len(WRITTEN) - 1))
        if i < len(WRITTEN) - 1:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    footer(d, "WA (SAQ)")
    P.enable_update_fields(d)
    p = os.path.join(OUT, f"WA (SAQ) - {C.SHORT_TITLE}.docx")
    d.save(p); print("  ", os.path.basename(p))


def build_wa_key():
    d = new_doc()
    cover(d, "Written Assessment (SAQ)", is_key=True)
    line(d, "Answer Key and Marking Guide", bold=True, size=14, color=BRAND, after=4)
    line(d, f"{C.TITLE}  ·  {C.COURSE_CODE}", size=10, color=GREY, after=4)
    line(d, f"Written Assessment (SAQ) — {len(WRITTEN)} questions · {WA_DURATION} · "
            f"open book", size=10, color=GREY, after=4)
    line(d, "Grading: the learner must answer all questions correctly to be assessed "
            "Competent (C). Answers need not match this key word for word — award "
            "Competent where the substance is covered.", size=10, color=GREY, after=12)
    for q in WRITTEN:
        key_answer(d, q["n"], q["a"], q["codes"], q["ref"])
    footer(d, "WA (SAQ) — ANSWER KEY")
    P.enable_update_fields(d)
    p = os.path.join(OUT, f"Answer to WA (SAQ) - {C.SHORT_TITLE}.docx")
    d.save(p); print("  ", os.path.basename(p))


def build_cs_paper():
    d = new_doc()
    cover(d, "Case Study (CS)")
    admin_page(d, "Case Study (CS)", CS_DURATION, len(CS_ITEMS),
               extra_instructions=[
                   "Each question is based on its own scenario. Read the scenario carefully "
                   "before answering."])
    line(d, "C: Scenarios and Questions", bold=True, size=12, after=10)
    for i, q in enumerate(CS_ITEMS):
        if i:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        for blk in q["scenario"].split("\n\n"):
            b = blk.strip()
            if not b:
                continue
            p = d.add_paragraph(); p.paragraph_format.space_after = Pt(6)
            P.md_runs(p, b, size=10.5)
        d.add_paragraph()
        question(d, q["n"], q["q"], q["codes"])
        answer_box(d, lines=12, trailing=(i < len(CS_ITEMS) - 1))
    footer(d, "Case Study (CS)")
    P.enable_update_fields(d)
    p = os.path.join(OUT, f"CS - {C.SHORT_TITLE}.docx")
    d.save(p); print("  ", os.path.basename(p))


def build_cs_key():
    d = new_doc()
    cover(d, "Case Study (CS)", is_key=True)
    line(d, "Answer Key and Marking Guide", bold=True, size=14, color=BRAND, after=4)
    line(d, f"{C.TITLE}  ·  {C.COURSE_CODE}", size=10, color=GREY, after=4)
    line(d, f"Case Study (CS) — {len(CS_ITEMS)} questions · {CS_DURATION} · open book",
         size=10, color=GREY, after=4)
    line(d, "Grading: the learner must answer all questions correctly to be assessed "
            "Competent (C). These are suggested answers — award Competent where the "
            "learner covers the substance and reasons from the scenario evidence.",
         size=10, color=GREY, after=12)
    for q in CS_ITEMS:
        key_answer(d, q["n"], q["a"], q["codes"], q["ref"])
    footer(d, "Case Study (CS) — ANSWER KEY")
    P.enable_update_fields(d)
    p = os.path.join(OUT, f"Answer to CS - {C.SHORT_TITLE}.docx")
    d.save(p); print("  ", os.path.basename(p))


if __name__ == "__main__":
    print("Coverage check:")
    check_coverage()
    print("Building assessment set ->", OUT)
    build_wa_paper()
    build_wa_key()
    build_cs_paper()
    build_cs_key()
    print("Done — 4 DOCX (2 question papers + 2 answer keys).")
