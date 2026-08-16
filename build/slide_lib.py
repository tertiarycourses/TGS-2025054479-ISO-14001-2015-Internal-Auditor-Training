#!/usr/bin/env python3
"""Component library for the ISO 14001:2015 Internal Auditor deck.

Ported from the WSQ compact-v2 reference (wsq-slides/reference/compact/build_slides.py)
and the measured grid spec in the tertiary-ppt-design skill. Every layout helper the
deck uses lives here so build_slides.py stays pure content.

House invariants enforced here:
  * 16:9, all-white, Arial only
  * footer on every numbered slide: course + code (left), copyright (centre), page (right)
  * titles auto-fit to ONE line (measured with PIL, not a character-count guess)
  * process/flow visuals use REAL shapes and REAL connectors with arrowheads
  * transitions applied in ONE pass at the end (fade content / push dividers)
"""
import os, math, json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, ".."))
ASSETS = os.path.join(REPO, "courseware", "assets")

# ---------------- palette ----------------
BLUE = RGBColor(0x1F, 0x6F, 0xEB); TEAL = RGBColor(0x10, 0xB9, 0x81)
AMBER = RGBColor(0xF5, 0x9E, 0x0B); VIOLET = RGBColor(0x7C, 0x3A, 0xED)
INK = RGBColor(0x16, 0x1B, 0x26); GREY = RGBColor(0x5B, 0x63, 0x72)
LIGHT = RGBColor(0xF5, 0xF8, 0xFC); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD7, 0xE0, 0xEA); RED = RGBColor(0xDC, 0x26, 0x26)
NAVY = RGBColor(0x0B, 0x12, 0x20); GREEN = RGBColor(0x16, 0x84, 0x5B)
MIST = RGBColor(0xE8, 0xF0, 0xFE); MINT = RGBColor(0xE8, 0xF7, 0xEE)
BLUSH = RGBColor(0xFE, 0xF2, 0xF2)
PALETTE = [BLUE, TEAL, VIOLET, AMBER, GREEN, RED]

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

PAGE = {"n": 0}
SLIDE_MAP = {}
_DIVIDERS = set()          # slide indexes that get `push` instead of `fade`

COURSE_FOOTER = {"text": "", "code": ""}


def set_footer(text, code):
    COURSE_FOOTER["text"] = text
    COURSE_FOOTER["code"] = code


# ---------------- primitives ----------------
def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, color, line=None, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, int(x), int(y), int(w), int(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def oval(s, x, y, w, h, color):
    sp = s.shapes.add_shape(MSO_SHAPE.OVAL, int(x), int(y), int(w), int(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=4):
    tb = s.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0
    for i, ln in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        for t, sz, col, bold in ln:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = bold
            r.font.color.rgb = col; r.font.name = "Arial"
    return tb


def bullets(s, x, y, w, h, items, size=18, color=INK, gap=10):
    tb = s.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        lvl = it[1] if isinstance(it, tuple) and len(it) > 1 else 0
        text = it[0] if isinstance(it, tuple) else it
        bold = bool(isinstance(it, tuple) and len(it) > 2 and it[2])
        r = p.add_run()
        r.text = ("•  " if lvl == 0 else "–  ") + text
        r.font.size = Pt(size if lvl == 0 else size - 2)
        r.font.color.rgb = color if lvl == 0 else GREY
        r.font.name = "Arial"; r.font.bold = bold
    return tb


# ---------------- title auto-fit (measured, not guessed) ----------------
_FONT_CACHE = {}
TITLE_BOX_IN = 11.75 * 0.95


def _pil_font(size_pt):
    if size_pt in _FONT_CACHE:
        return _FONT_CACHE[size_pt]
    try:
        from PIL import ImageFont
        for path in ("/Library/Fonts/Arial Bold.ttf",
                     "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
                     "/Library/Fonts/Arial.ttf"):
            if os.path.exists(path):
                f = ImageFont.truetype(path, int(size_pt * 4))
                _FONT_CACHE[size_pt] = f
                return f
    except Exception:
        pass
    _FONT_CACHE[size_pt] = None
    return None


def _fit_title_size(text):
    """Largest of 28/26/24/22/21 pt that fits on ONE line in the title box."""
    for size in (28, 26, 24, 22, 21):
        f = _pil_font(size)
        if f is None:
            return 28 if len(text) <= 46 else (24 if len(text) <= 58 else 21)
        w_px = f.getbbox(text)[2]
        w_in = w_px / (4 * 72.0)
        if w_in <= TITLE_BOX_IN:
            return size
    return 21


def _ellipsis(text, limit):
    t = " ".join(str(text).split())
    if len(t) <= limit:
        return t
    cut = t[:limit]; sp = cut.rfind(" ")
    if sp > limit * 0.55:
        cut = cut[:sp]
    return cut.rstrip(" ,.;:-—(") + "…"


# ---------------- chrome ----------------
def footer(s):
    PAGE["n"] += 1
    txt(s, Inches(0.4), Inches(7.04), Inches(7.2), Inches(0.34),
        [[(f"{COURSE_FOOTER['text']}  ·  {COURSE_FOOTER['code']}", 8.5, GREY, False)]])
    txt(s, Inches(4.9), Inches(7.04), Inches(3.6), Inches(0.34),
        [[("© 2026 Tertiary Infotech Academy Pte Ltd", 8.5, GREY, False)]],
        align=PP_ALIGN.CENTER)
    txt(s, Inches(12.35), Inches(7.04), Inches(0.6), Inches(0.34),
        [[(str(PAGE["n"]), 8.5, GREY, False)]], align=PP_ALIGN.RIGHT)


def head(s, title, kicker=None, kcolor=BLUE):
    rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, Inches(0.22), Inches(1.48), kcolor)
    if kicker:
        txt(s, Inches(0.72), Inches(0.34), Inches(11.85), Inches(0.34),
            [[(kicker.upper(), 12, kcolor, True)]])
    txt(s, Inches(0.72), Inches(0.68), Inches(11.85), Inches(0.72),
        [[(title, _fit_title_size(title), INK, True)]])
    rect(s, Inches(0.72), Inches(1.50), Inches(11.85), Inches(0.012), LINE)
    return s


def mark(key):
    SLIDE_MAP[key] = PAGE["n"] + 1


def _asset(name):
    p = os.path.join(ASSETS, name)
    return p if os.path.exists(p) else None


# ---------------- components ----------------
def cover(C, badge_top="ISO", badge_sub="14001:2015"):
    s = slide(); rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, SW, Inches(0.22), BLUE)
    rect(s, 0, Inches(7.28), SW, Inches(0.22), TEAL)
    logo = _asset("tertiary-infotech-logo.png")
    if logo:
        s.shapes.add_picture(logo, Inches(0.85), Inches(0.62), height=Inches(1.0))
    rect(s, Inches(10.75), Inches(0.64), Inches(1.8), Inches(1.0), TEAL)
    txt(s, Inches(10.75), Inches(0.76), Inches(1.8), Inches(0.45),
        [[(badge_top, 21, WHITE, True)]], align=PP_ALIGN.CENTER)
    txt(s, Inches(10.75), Inches(1.22), Inches(1.8), Inches(0.36),
        [[(badge_sub, 10, WHITE, True)]], align=PP_ALIGN.CENTER)
    txt(s, Inches(0.88), Inches(2.22), Inches(12), Inches(0.5),
        [[("COURSE SLIDES  ·  WSQ", 15, BLUE, True)]])
    txt(s, Inches(0.88), Inches(2.72), Inches(11.6), Inches(1.9),
        [[(C.TITLE, 40, INK, True)]])
    rect(s, Inches(0.90), Inches(4.42), Inches(2.4), Inches(0.06), TEAL)
    txt(s, Inches(0.88), Inches(4.76), Inches(12), Inches(1.5),
        [[(f"WSQ Course Code: {C.COURSE_CODE}  ·  {C.DAYS}-Day Instructor-Led Training", 15, GREY, False)],
         [(f"Skills Framework TSC: {C.TSC_TITLE} ({C.TSC_CODE})", 13, GREY, False)],
         [(f"Conducted by {C.ORG}  ·  {C.UEN}", 13, GREY, False)]], space=6)
    txt(s, Inches(0.88), Inches(6.42), Inches(12), Inches(0.36),
        [[(f"Version {C.VERSION}  ·  {C.VERSION_DATE}", 12, GREY, False)]])
    txt(s, Inches(0.88), Inches(6.80), Inches(12), Inches(0.32),
        [[("© 2026 Tertiary Infotech Academy Pte Ltd. All rights reserved.  ·  "
           "www.tertiarycourses.com.sg", 9.5, GREY, False)]])
    PAGE["n"] += 1


def section(kicker, title, num, sub=""):
    s = slide(); rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, Inches(0.22), SH, BLUE)
    rect(s, Inches(0.72), Inches(1.90), Inches(11.85), Inches(3.60), LIGHT)
    rect(s, Inches(0.72), Inches(1.90), Inches(0.14), Inches(3.60), TEAL)
    txt(s, Inches(1.40), Inches(2.30), Inches(10.8), Inches(0.45),
        [[(kicker.upper(), 15, BLUE, True)]])
    txt(s, Inches(1.40), Inches(2.80), Inches(10.9), Inches(1.15),
        [[(title, 34, INK, True)]])
    if sub:
        txt(s, Inches(1.40), Inches(3.98), Inches(10.7), Inches(1.3),
            [[(sub, 14.5, GREY, False)]])
    txt(s, Inches(9.9), Inches(5.62), Inches(2.6), Inches(1.3),
        [[(num, 60, LINE, True)]], align=PP_ALIGN.RIGHT)
    _DIVIDERS.add(len(prs.slides) - 1)
    footer(s); return s


def content(title, items, kicker=None, size=18, kcolor=BLUE):
    s = head(slide(), title, kicker, kcolor)
    bullets(s, Inches(0.72), Inches(1.90), Inches(11.85), Inches(4.9), items, size=size)
    footer(s); return s


def tile_grid(title, items, kicker=None, cols=2, size=14, accent=BLUE, icons=None):
    """Grid of light panels: accent stripe + numbered badge + bold title + caption."""
    s = head(slide(), title, kicker, kcolor=accent)
    n = len(items); rows = math.ceil(n / cols)
    X0 = Inches(0.72); Y0 = Inches(1.88); TOTW = Inches(11.85); AREAH = Inches(4.85)
    gx = Inches(0.28); gy = Inches(0.22)
    cw = int((TOTW - gx * (cols - 1)) / cols)
    ch = int((AREAH - gy * (rows - 1)) / rows)
    bd = Inches(0.52)
    for i, it in enumerate(items):
        r, c = i // cols, i % cols
        x = int(X0 + (cw + gx) * c); y = int(Y0 + (ch + gy) * r)
        col = PALETTE[i % len(PALETTE)]
        rect(s, x, y, cw, ch, LIGHT)
        rect(s, x, y, Inches(0.09), ch, col)
        oval(s, x + Inches(0.26), int(y + ch / 2 - bd / 2), bd, bd, col)
        ic = icons[i] if icons else str(i + 1)
        txt(s, x + Inches(0.26), int(y + ch / 2 - bd / 2), bd, bd,
            [[(ic, 16, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tx = x + Inches(0.96); tw = cw - Inches(1.20)
        if isinstance(it, tuple):
            txt(s, tx, int(y + Inches(0.12)), tw, int(ch - Inches(0.18)),
                [[(it[0], size + 1, INK, True)], [(it[1], size - 2, GREY, False)]],
                anchor=MSO_ANCHOR.MIDDLE, space=3)
        else:
            txt(s, tx, int(y + Inches(0.10)), tw, int(ch - Inches(0.16)),
                [[(it, size, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s


def cards(title, items, kicker=None, accent=BLUE, note=None):
    """3 or 4 across concept cards — the five-layer card of the measured grid spec."""
    s = head(slide(), title, kicker, kcolor=accent)
    n = len(items)
    if n <= 2:   xs, cw = [Inches(0.72), Inches(6.72)], Inches(5.85)
    elif n == 3: xs, cw = [Inches(0.72), Inches(4.71), Inches(8.70)], Inches(3.71)
    else:        xs, cw = [Inches(0.72), Inches(3.72), Inches(6.72), Inches(9.72)], Inches(2.72)
    y = Inches(1.95); ch = Inches(2.55) if note else Inches(3.30)
    for i, (t, body) in enumerate(items[:4]):
        x = xs[i]; col = PALETTE[i % len(PALETTE)]
        rect(s, x, y, cw, ch, LIGHT)
        rect(s, x, y, Inches(0.09), ch, col)
        bd = Inches(0.46)
        oval(s, x + Inches(0.24), y + Inches(0.20), bd, bd, col)
        txt(s, x + Inches(0.24), y + Inches(0.20), bd, bd, [[(str(i + 1), 15, WHITE, True)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x + Inches(0.82), y + Inches(0.14), cw - Inches(1.0), Inches(0.56),
            [[(t, 14, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x + Inches(0.28), y + Inches(0.78), cw - Inches(0.54), ch - Inches(0.92),
            [[(body, 11, INK, False)]])
    if note:
        by = Inches(4.72)
        rect(s, Inches(0.72), by, Inches(11.85), Inches(1.42), LIGHT)
        txt(s, Inches(1.0), by + Inches(0.13), Inches(11.3), Inches(0.32),
            [[("THE POINT", 12, accent, True)]])
        txt(s, Inches(1.0), by + Inches(0.48), Inches(11.3), Inches(0.86),
            [[(note, 12, INK, False)]])
    footer(s); return s


def two_panel(title, lhead, litems, rhead, ritems, kicker=None,
              lcolor=BLUE, rcolor=TEAL, note=None):
    s = head(slide(), title, kicker, kcolor=lcolor)
    bh = Inches(4.15) if note else Inches(4.85)
    y = Inches(1.90)
    for x, w, hd, items, col in ((Inches(0.72), Inches(5.80), lhead, litems, lcolor),
                                 (Inches(6.77), Inches(5.80), rhead, ritems, rcolor)):
        rect(s, x, y, w, bh, LIGHT)
        rect(s, x, y, w, Inches(0.42), col)
        txt(s, x + Inches(0.24), y, w - Inches(0.4), Inches(0.42),
            [[(hd, 14, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
        bullets(s, x + Inches(0.26), y + Inches(0.60), w - Inches(0.52),
                bh - Inches(0.75), items, size=13, gap=7)
    if note:
        rect(s, Inches(0.72), Inches(6.18), Inches(11.85), Inches(0.72), MIST)
        txt(s, Inches(1.0), Inches(6.18), Inches(11.3), Inches(0.72),
            [[(note, 12, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s


def split_note(title, lhead, litems, rhead, ritems, kicker=None, note=None):
    """Do / Don't panels — red blush vs green mint."""
    s = head(slide(), title, kicker, kcolor=GREEN)
    bh = Inches(4.15) if note else Inches(4.85)
    y = Inches(1.90)
    for x, w, hd, items, col, fill in (
            (Inches(0.72), Inches(5.80), lhead, litems, RED, BLUSH),
            (Inches(6.77), Inches(5.80), rhead, ritems, GREEN, MINT)):
        rect(s, x, y, w, bh, fill)
        rect(s, x, y, w, Inches(0.42), col)
        txt(s, x + Inches(0.24), y, w - Inches(0.4), Inches(0.42),
            [[(hd, 14, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
        bullets(s, x + Inches(0.26), y + Inches(0.60), w - Inches(0.52),
                bh - Inches(0.75), items, size=13, gap=7)
    if note:
        rect(s, Inches(0.72), Inches(6.18), Inches(11.85), Inches(0.72), LIGHT)
        txt(s, Inches(1.0), Inches(6.18), Inches(11.3), Inches(0.72),
            [[(note, 12, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s


def process_map(title, stages, kicker=None, accent=BLUE, synthesis=None):
    """REAL rounded-rect stages joined by REAL connectors with arrowheads."""
    s = head(slide(), title, kicker, kcolor=accent)
    n = len(stages)
    X0 = Inches(0.72); TOTW = Inches(11.85); gap = Inches(0.34)
    cw = int((TOTW - gap * (n - 1)) / n)
    y = Inches(2.30); ch = Inches(2.60) if synthesis else Inches(3.15)
    bd = Inches(0.78)
    for i, st in enumerate(stages):
        label, cap = (st if isinstance(st, tuple) else (st, ""))
        x = int(X0 + (cw + gap) * i)
        col = PALETTE[i % len(PALETTE)] if accent is None else accent
        sp = rect(s, x, y, cw, ch, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(s, x, y, cw, Inches(0.10), col)
        oval(s, int(x + cw / 2 - bd / 2), int(y + Inches(0.38)), bd, bd, col)
        txt(s, int(x + cw / 2 - bd / 2), int(y + Inches(0.38)), bd, bd,
            [[(str(i + 1), 26, WHITE, True)]], align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x + Inches(0.12), int(y + Inches(1.32)), cw - Inches(0.24), Inches(0.62),
            [[(label, 12.5, INK, True)]], align=PP_ALIGN.CENTER)
        if cap:
            txt(s, x + Inches(0.12), int(y + Inches(1.92)), cw - Inches(0.24),
                ch - Inches(2.02), [[(cap, 10, GREY, False)]], align=PP_ALIGN.CENTER)
        if i < n - 1:
            cx1 = int(x + cw + Inches(0.05)); cx2 = int(x + cw + gap - Inches(0.05))
            cy = int(y + ch / 2)
            cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cx1, cy, cx2, cy)
            cn.line.color.rgb = col; cn.line.width = Pt(2.25)
            ln = cn.line._get_or_add_ln()
            te = etree.SubElement(ln, qn("a:tailEnd"))
            te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    if synthesis:
        by = Inches(5.18)
        rect(s, Inches(0.72), by, Inches(11.85), Inches(1.42), LIGHT)
        txt(s, Inches(1.0), by + Inches(0.13), Inches(11.3), Inches(0.32),
            [[("THE POINT", 12, accent, True)]])
        txt(s, Inches(1.0), by + Inches(0.48), Inches(11.3), Inches(0.86),
            [[(synthesis, 12, INK, False)]])
    footer(s); return s


def flow_h(title, steps, kicker=None, color=BLUE, note=None):
    """Numbered chips joined by chevrons — the Assessment Flow shape."""
    s = head(slide(), title, kicker, kcolor=color)
    n = len(steps); X0 = Inches(0.72); TOTW = Inches(11.85); gap = Inches(0.34)
    cw = int((TOTW - gap * (n - 1)) / n)
    y = Inches(2.45); ch = Inches(3.15); bd = Inches(0.82)
    for i, st in enumerate(steps):
        x = int(X0 + (cw + gap) * i)
        rect(s, x, y, cw, ch, LIGHT)
        rect(s, x, y, cw, Inches(0.10), color)
        oval(s, int(x + cw / 2 - bd / 2), int(y + Inches(0.42)), bd, bd, color)
        txt(s, int(x + cw / 2 - bd / 2), int(y + Inches(0.42)), bd, bd,
            [[(str(i + 1), 28, WHITE, True)]], align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x + Inches(0.14), int(y + Inches(1.52)), cw - Inches(0.28),
            int(ch - Inches(1.66)), [[(st, 12, INK, False)]], align=PP_ALIGN.CENTER)
        if i < n - 1:
            ch_sp = s.shapes.add_shape(MSO_SHAPE.CHEVRON,
                                       int(x + cw + Inches(0.02)),
                                       int(y + ch / 2 - Inches(0.20)),
                                       int(gap - Inches(0.04)), Inches(0.40))
            ch_sp.fill.solid(); ch_sp.fill.fore_color.rgb = color
            ch_sp.line.fill.background(); ch_sp.shadow.inherit = False
    if note:
        txt(s, Inches(0.72), Inches(5.95), Inches(11.85), Inches(0.7),
            [[(note, 12.5, GREY, False)]], align=PP_ALIGN.CENTER)
    footer(s); return s


def table_slide(title, headers, rows, kicker=None, accent=BLUE,
                widths=None, note=None, fsize=12):
    s = head(slide(), title, kicker, kcolor=accent)
    ncol = len(headers); X0 = Inches(0.72); TOTW = Inches(11.85)
    ws = [int(TOTW * w) for w in widths] if widths else [int(TOTW / ncol)] * ncol
    area_h = Inches(4.32) if note else Inches(4.88)
    nrow = len(rows) + 1
    rh = int(area_h / nrow); y0 = Inches(1.90)
    x = X0
    for j, h in enumerate(headers):
        rect(s, x, y0, ws[j], rh, accent)
        txt(s, x + Inches(0.12), y0, ws[j] - Inches(0.22), rh,
            [[(h, fsize, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
        x += ws[j]
    for i, row in enumerate(rows):
        y = int(y0 + rh * (i + 1)); x = X0
        fill = LIGHT if i % 2 == 0 else WHITE
        for j, cell in enumerate(row):
            rect(s, x, y, ws[j], rh, fill, line=LINE)
            txt(s, x + Inches(0.12), y, ws[j] - Inches(0.22), rh,
                [[(str(cell), fsize - 1, INK, j == 0)]], anchor=MSO_ANCHOR.MIDDLE)
            x += ws[j]
    if note:
        rect(s, Inches(0.72), Inches(6.30), Inches(11.85), Inches(0.62), MIST)
        txt(s, Inches(1.0), Inches(6.30), Inches(11.3), Inches(0.62),
            [[(note, 11.5, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s


def img_points(title, image, points, kicker=None, accent=BLUE, img_w=7.0, note=None):
    """Diagram left + colour-coded takeaway tiles right — the workhorse concept slide."""
    s = head(slide(), title, kicker, kcolor=accent)
    p = _asset(image)
    area_h = Inches(4.35) if note else Inches(4.85)
    rx = Inches(6.6)
    if p:
        from PIL import Image
        iw, ih = Image.open(p).size
        w = Inches(img_w); h = int(w * ih / iw)
        if h > area_h:
            h = int(area_h); w = int(h * iw / ih)
        y = int(Inches(1.90) + (area_h - h) / 2)
        s.shapes.add_picture(p, Inches(0.72), y, width=int(w), height=int(h))
        rx = int(Inches(0.72) + w + Inches(0.32))
    rw = int(Inches(12.57) - rx)
    n = len(points); gy = Inches(0.18)
    th = int((area_h - gy * (n - 1)) / n)
    for i, (t, sub) in enumerate(points):
        y = int(Inches(1.90) + (th + gy) * i)
        col = PALETTE[i % len(PALETTE)]
        rect(s, rx, y, rw, th, LIGHT)
        rect(s, rx, y, Inches(0.09), th, col)
        txt(s, rx + Inches(0.26), y, rw - Inches(0.46), th,
            [[(t, 13, col, True)], [(sub, 11, INK, False)]],
            anchor=MSO_ANCHOR.MIDDLE, space=2)
    if note:
        rect(s, Inches(0.72), Inches(6.38), Inches(11.85), Inches(0.62), MIST)
        txt(s, Inches(1.0), Inches(6.38), Inches(11.3), Inches(0.62),
            [[(note, 11.5, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s


def img_full(title, image, kicker=None, accent=BLUE, caption=None):
    s = head(slide(), title, kicker, kcolor=accent)
    p = _asset(image)
    area_h = Inches(4.30) if caption else Inches(4.85)
    if p:
        from PIL import Image
        iw, ih = Image.open(p).size
        w = Inches(11.5); h = int(w * ih / iw)
        if h > area_h:
            h = int(area_h); w = int(h * iw / ih)
        x = int((SW - w) / 2); y = int(Inches(1.88) + (area_h - h) / 2)
        s.shapes.add_picture(p, x, y, width=int(w), height=int(h))
    if caption:
        rect(s, Inches(0.72), Inches(6.28), Inches(11.85), Inches(0.72), LIGHT)
        txt(s, Inches(1.0), Inches(6.28), Inches(11.3), Inches(0.72),
            [[(caption, 12, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s


def stack(title, items, kicker=None, accent=BLUE, note=None):
    """Vertical numbered sequence — badge + bold title + grey caption per row."""
    s = head(slide(), title, kicker, kcolor=accent)
    n = len(items)
    area_h = Inches(4.35) if note else Inches(4.88)
    gy = Inches(0.16); rh = int((area_h - gy * (n - 1)) / n)
    for i, it in enumerate(items):
        t, sub = it if isinstance(it, tuple) else (it, "")
        y = int(Inches(1.90) + (rh + gy) * i)
        col = PALETTE[i % len(PALETTE)]
        rect(s, Inches(0.72), y, Inches(11.85), rh, LIGHT)
        rect(s, Inches(0.72), y, Inches(0.09), rh, col)
        bd = min(Inches(0.50), rh - Inches(0.10))
        oval(s, Inches(0.98), int(y + rh / 2 - bd / 2), bd, bd, col)
        txt(s, Inches(0.98), int(y + rh / 2 - bd / 2), bd, bd,
            [[(str(i + 1), 14, WHITE, True)]], align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
        if sub:
            txt(s, Inches(1.68), y, Inches(10.7), rh,
                [[(t, 13, INK, True)], [(sub, 11, GREY, False)]],
                anchor=MSO_ANCHOR.MIDDLE, space=2)
        else:
            txt(s, Inches(1.68), y, Inches(10.7), rh, [[(t, 13, INK, False)]],
                anchor=MSO_ANCHOR.MIDDLE)
    if note:
        rect(s, Inches(0.72), Inches(6.40), Inches(11.85), Inches(0.62), MIST)
        txt(s, Inches(1.0), Inches(6.40), Inches(11.3), Inches(0.62),
            [[(note, 11.5, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s


def big_statement(line1, line2, kicker, color=BLUE):
    s = slide(); rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, Inches(0.22), SH, color)
    txt(s, Inches(1.05), Inches(2.20), Inches(11), Inches(0.45),
        [[(kicker.upper(), 15, color, True)]])
    txt(s, Inches(1.05), Inches(2.75), Inches(11.3), Inches(2.3),
        [[(line1, 36, INK, True)]])
    if line2:
        txt(s, Inches(1.07), Inches(4.85), Inches(11), Inches(1.2),
            [[(line2, 18, GREY, False)]])
    footer(s); return s


def trainer_slide(kicker, name, role, rows, initials, accent=BLUE):
    s = head(slide(), "About the Trainer", kicker, kcolor=accent)
    lx = Inches(0.72); lw = Inches(3.60)
    rect(s, lx, Inches(1.90), lw, Inches(4.85), LIGHT)
    rect(s, lx, Inches(1.90), lw, Inches(0.12), accent)
    bd = Inches(1.65); ax = int(lx + (lw - bd) / 2)
    oval(s, ax, Inches(2.45), bd, bd, accent)
    txt(s, ax, Inches(2.45), bd, bd, [[(initials, 42, WHITE, True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, lx + Inches(0.15), Inches(4.45), lw - Inches(0.3), Inches(0.6),
        [[(name, 19, INK, True)]], align=PP_ALIGN.CENTER)
    txt(s, lx + Inches(0.15), Inches(5.10), lw - Inches(0.3), Inches(1.4),
        [[(ln, 12, GREY, False)] for ln in role.split("\n")],
        align=PP_ALIGN.CENTER, space=2)
    rx = Inches(4.62); rw = Inches(7.95); ry = Inches(1.90); rh = Inches(4.85)
    n = len(rows); gy = Inches(0.16); th = int((rh - gy * (n - 1)) / n)
    for i, (label, val) in enumerate(rows):
        y = int(ry + (th + gy) * i); col = PALETTE[i % len(PALETTE)]
        rect(s, rx, y, rw, th, LIGHT)
        rect(s, rx, y, Inches(0.09), th, col)
        vruns = [(val, 12.5, INK, False)] if val else \
                [("________________________________________________", 12, LINE, False)]
        txt(s, rx + Inches(0.30), y, rw - Inches(0.56), th,
            [[(label.upper(), 10, col, True)], vruns],
            anchor=MSO_ANCHOR.MIDDLE, space=3)
    footer(s); return s


def brk(kind, dur, color=AMBER):
    s = slide(); rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, SW, Inches(0.22), color)
    rect(s, 0, Inches(7.28), SW, Inches(0.22), color)
    rect(s, Inches(5.4), Inches(2.35), Inches(2.53), Inches(0.10), color)
    txt(s, 0, Inches(2.75), SW, Inches(1.2), [[(kind, 44, INK, True)]],
        align=PP_ALIGN.CENTER)
    txt(s, 0, Inches(4.05), SW, Inches(0.8), [[(dur, 20, color, True)]],
        align=PP_ALIGN.CENTER)
    _DIVIDERS.add(len(prs.slides) - 1)
    PAGE["n"] += 1


def activity_slide(a):
    """ONE workflow slide per activity: scenario + 5-chip strip + deliverable + tools."""
    s = head(slide(), f"Activity {a['num']} — {_ellipsis(a['title'], 46)}",
             kicker=f"TOPIC {a['topic']} · HANDS-ON ACTIVITY", kcolor=TEAL)
    rect(s, Inches(10.30), Inches(0.32), Inches(2.27), Inches(0.42), TEAL)
    txt(s, Inches(10.30), Inches(0.32), Inches(2.27), Inches(0.42),
        [[(f"ACTIVITY {a['num']}  ·  {a['minutes']} MIN", 12, WHITE, True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(0.72), Inches(1.88), Inches(11.85), Inches(1.15),
        [[(a["desc"], 13.5, INK, False)]])
    steps = a["flow"]; n = len(steps)
    X0 = Inches(0.72); TOTW = Inches(11.85); gap = Inches(0.26)
    cw = int((TOTW - gap * (n - 1)) / n)
    y = Inches(3.16); ch = Inches(2.02); bd = Inches(0.54)
    for i, st in enumerate(steps):
        x = int(X0 + (cw + gap) * i)
        rect(s, x, y, cw, ch, LIGHT)
        rect(s, x, y, cw, Inches(0.08), TEAL)
        oval(s, int(x + cw / 2 - bd / 2), int(y + Inches(0.22)), bd, bd, TEAL)
        txt(s, int(x + cw / 2 - bd / 2), int(y + Inches(0.22)), bd, bd,
            [[(str(i + 1), 19, WHITE, True)]], align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x + Inches(0.10), int(y + Inches(0.90)), cw - Inches(0.20),
            int(ch - Inches(0.98)), [[(st, 10.5, INK, False)]], align=PP_ALIGN.CENTER)
        if i < n - 1:
            c = s.shapes.add_shape(MSO_SHAPE.CHEVRON,
                                   int(x + cw + Inches(0.01)),
                                   int(y + ch / 2 - Inches(0.16)),
                                   int(gap - Inches(0.02)), Inches(0.32))
            c.fill.solid(); c.fill.fore_color.rgb = TEAL
            c.line.fill.background(); c.shadow.inherit = False
    rect(s, Inches(0.72), Inches(5.40), Inches(11.85), Inches(0.86), MINT)
    txt(s, Inches(1.0), Inches(5.40), Inches(11.3), Inches(0.86),
        [[("YOU'LL PRODUCE   ", 11.5, GREEN, True), (a["build"], 12, INK, False)]],
        anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(0.72), Inches(6.38), Inches(11.85), Inches(0.5),
        [[("Materials:  ", 11, GREY, True), (a["services"], 11, GREY, False),
          (f"     ·     Full step-by-step guide: Learner Guide, Activity {a['num']}",
           11, GREY, False)]])
    footer(s); return s


def verify_slide(a):
    """Green acceptance-criterion slide with a troubleshooting band."""
    s = head(slide(), f"Activity {a['num']} — Check Your Work",
             kicker=f"TOPIC {a['topic']} · SELF-CHECK", kcolor=GREEN)
    rect(s, Inches(0.72), Inches(1.90), Inches(11.85), Inches(1.72), MINT)
    txt(s, Inches(1.0), Inches(2.02), Inches(11.3), Inches(0.34),
        [[("YOU ARE DONE WHEN", 12, GREEN, True)]])
    txt(s, Inches(1.0), Inches(2.40), Inches(11.3), Inches(1.10),
        [[(a["test"], 13, INK, False)]])
    rect(s, Inches(0.72), Inches(3.78), Inches(11.85), Inches(2.98), LIGHT)
    txt(s, Inches(1.0), Inches(3.92), Inches(11.3), Inches(0.34),
        [[("IF YOU ARE STUCK — WHAT USUALLY GOES WRONG", 12, AMBER, True)]])
    bullets(s, Inches(1.0), Inches(4.32), Inches(11.3), Inches(2.34),
            a.get("pitfalls", []), size=12, gap=8)
    footer(s); return s


def lms_slide(course_title):
    """Download Course Material — browser-card mock + numbered how-to tiles."""
    s = head(slide(), "Download Course Material", kicker="COURSE PORTAL · LMS/TMS",
             kcolor=BLUE)
    bx, by, bw, bh = Inches(0.72), Inches(1.95), Inches(6.05), Inches(4.55)
    rect(s, bx, by, bw, bh, WHITE, line=LINE)
    rect(s, bx, by, bw, Inches(0.50), RGBColor(0xEE, 0xF2, 0xF8))
    for i, c in enumerate([RED, AMBER, TEAL]):
        oval(s, int(bx + Inches(0.18) + Inches(0.28) * i), int(by + Inches(0.16)),
             Inches(0.17), Inches(0.17), c)
    rect(s, int(bx + Inches(1.15)), int(by + Inches(0.09)), int(bw - Inches(1.45)),
         Inches(0.33), WHITE, line=LINE)
    txt(s, int(bx + Inches(1.30)), int(by + Inches(0.09)), int(bw - Inches(1.75)),
        Inches(0.33), [[("https://lms-tms.tertiaryinfotech.com", 12, BLUE, True)]],
        anchor=MSO_ANCHOR.MIDDLE)
    rect(s, int(bx + Inches(0.24)), int(by + Inches(0.72)), int(bw - Inches(0.48)),
         Inches(0.72), BLUE)
    txt(s, int(bx + Inches(0.44)), int(by + Inches(0.72)), int(bw - Inches(0.88)),
        Inches(0.72),
        [[("Tertiary Infotech LMS/TMS", 14, WHITE, True)],
         [("My Courses  ·  Attendance  ·  Assessment", 10, RGBColor(0xB3, 0xD4, 0xFF), False)]],
        anchor=MSO_ANCHOR.MIDDLE, space=2)
    rows = [(_ellipsis(course_title, 44), True), ("Trainer Slides (PDF)   ⬇", False),
            ("Learner Guide (PDF)   ⬇", False), ("Assessment Submission   ↗", False)]
    for i, (rt, hd) in enumerate(rows):
        yy = int(by + Inches(1.62) + Inches(0.66) * i)
        rect(s, int(bx + Inches(0.24)), yy, int(bw - Inches(0.48)), Inches(0.54),
             MIST if hd else LIGHT)
        txt(s, int(bx + Inches(0.44)), yy, int(bw - Inches(0.88)), Inches(0.54),
            [[(rt, 11.5, BLUE if hd else INK, hd)]], anchor=MSO_ANCHOR.MIDDLE)
    steps = [("Sign in", "lms-tms.tertiaryinfotech.com — log in with your registered email."),
             ("Open your course", "Select this course under My Courses."),
             ("Download materials", "Slides, Learner Guide and the activity briefs — your "
                                    "open-book references."),
             ("Submit & attendance", "Submit your assessment answers and scan the digital "
                                     "attendance QR here.")]
    rx = Inches(7.15); rw = Inches(5.42); gy = Inches(0.19)
    th = int((Inches(4.55) - gy * 3) / 4)
    for i, (t1, t2) in enumerate(steps):
        y = int(Inches(1.95) + (th + gy) * i); col = PALETTE[i % 4]
        rect(s, rx, y, rw, th, LIGHT)
        rect(s, rx, y, Inches(0.09), th, col)
        bd = Inches(0.48)
        oval(s, rx + Inches(0.20), int(y + th / 2 - bd / 2), bd, bd, col)
        txt(s, rx + Inches(0.20), int(y + th / 2 - bd / 2), bd, bd,
            [[(str(i + 1), 15, WHITE, True)]], align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, rx + Inches(0.86), y, rw - Inches(1.06), th,
            [[(t1, 13, col, True)], [(t2, 10.5, INK, False)]],
            anchor=MSO_ANCHOR.MIDDLE, space=2)
    txt(s, Inches(0.72), Inches(6.58), Inches(11.85), Inches(0.4),
        [[("All course material is downloaded from the LMS/TMS portal — keep it handy: "
           "both assessments are open book.", 12, GREY, False)]], align=PP_ALIGN.CENTER)
    footer(s); return s


# ---------------- motion ----------------
def _transition(s, kind="fade", speed="fast"):
    sld = s._element
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    tr = etree.SubElement(sld, qn("p:transition"))
    tr.set("spd", speed); tr.set("advClick", "1")
    if kind == "fade":
        etree.SubElement(tr, qn("p:fade"))
    elif kind == "push":
        etree.SubElement(tr, qn("p:push")).set("dir", "l")
    sld.append(tr)


def apply_motion():
    """One pass at the very end: dividers push (med), everything else fade (fast)."""
    for i, s in enumerate(prs.slides):
        if i in _DIVIDERS:
            _transition(s, "push", "med")
        else:
            _transition(s, "fade", "fast")


def save(path, map_path=None):
    apply_motion()
    prs.save(path)
    if map_path:
        with open(map_path, "w") as f:
            json.dump(SLIDE_MAP, f, indent=1)
    return PAGE["n"]
